from __future__ import annotations

import html
import os
import subprocess
import sys
import traceback
import uuid
from datetime import datetime
from pathlib import Path

from PyQt6.QtCore import QByteArray, QRectF, Qt, QTimer, QUrl
from PyQt6.QtGui import QColor, QDesktopServices, QFont, QGuiApplication, QPainter, QPixmap
from PyQt6.QtWidgets import (
    QAbstractItemView, QApplication, QCheckBox, QColorDialog, QComboBox, QDoubleSpinBox, QFileDialog,
    QFrame, QGridLayout, QHBoxLayout, QHeaderView, QInputDialog, QLabel, QLineEdit, QMainWindow,
    QMessageBox, QProgressBar, QPushButton, QScrollArea, QSizePolicy, QSpinBox, QSplitter,
    QStackedWidget, QTableWidget, QTableWidgetItem, QTextBrowser, QTextEdit, QVBoxLayout, QWidget,
)

from storage import APP_DIR, APP_NAME, APP_VERSION, LOG_PATH, ConfigStore, SecretStore, VoiceStore
from theme import ACCENTS, LOG_COLORS, PROVIDER_ACCENT, STATUS_COLORS, build_stylesheet, make_arrow_icons
from media import VIDEO_SIZES, find_ffmpeg, make_video, probe_duration
from providers import LANGUAGES, MODELS, SPEED_RANGE, build_provider, language_label
from utils import fmt_duration, safe_filename, sample_hint, split_text, validate_voice_sample
from widgets import Card, ChoiceRow, ResponsiveRow, StatTile, button, label
from workers import BatchWorker, FuncWorker, produce_outputs
import updater
from app_info import GITHUB_REPO

FORMAT_OPTIONS = [("mp3", "🎵  MP3", "blue"), ("mp4", "🎬  MP4", "pink"), ("both", "🎵+🎬  Cả hai", "violet")]
PREVIEW_TEXT = {
    "vi": "Xin chào! Đây là giọng đọc thử của tôi. Chúc bạn một ngày thật vui vẻ.",
    "en": "Hello! This is a quick preview of my voice. Have a wonderful day.",
}
STATUS_TEXT = {"pending": "⏺ Chờ", "running": "⏳ Đang chạy", "ok": "✅ Xong", "error": "❌ Lỗi",
               "skipped": "⏭ Bỏ qua", "stopped": "⏹ Đã dừng"}
NAV = [
    ("🧬", "Clone giọng", "violet"),
    ("🗣", "Đọc văn bản", "blue"),
    ("📊", "Hàng loạt Excel", "green"),
    ("🎬", "Video MP4", "pink"),
    ("⚙", "Cài đặt API", "amber"),
    ("🔄", "Cập nhật", "teal"),
    ("📖", "Hướng dẫn", "cyan"),
]
PAGE_CLONE, PAGE_TTS, PAGE_BATCH, PAGE_VIDEO, PAGE_SETTINGS, PAGE_UPDATE, PAGE_GUIDE = range(7)
UPDATE_EVERY_MS = 6 * 3600 * 1000


def resource_path(rel: str) -> str:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return str(base / rel)


def unique_stem(folder: str, stem: str, exts: list[str]) -> str:
    stem = safe_filename(stem, "tts_output")
    cand, i = stem, 2
    while any((Path(folder) / f"{cand}{e}").exists() for e in exts):
        cand = f"{stem}_{i}"
        i += 1
    return cand


def exts_for(fmt: str) -> list[str]:
    return {"mp3": [".mp3"], "mp4": [".mp4"], "both": [".mp3", ".mp4"]}[fmt]


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"{APP_NAME} {APP_VERSION}")
        self.config_store = ConfigStore()
        self.voice_store = VoiceStore()
        self.secret_store = SecretStore()
        self.config = self.config_store.load()
        self.mode = self.config.get("theme", "dark") if self.config.get("theme") in ("dark", "light") else "dark"

        self.workers: set = set()
        self.responsive: list[ResponsiveRow] = []
        self.clone_worker = self.tts_worker = self.batch_worker = None
        self.batch_path = ""
        self.batch_tasks: list[dict] = []
        self.batch_state: list[str] = []
        self.batch_map: list[int] = []
        self.last_outputs: list[str] = []
        self.latest_release = None
        self.update_worker = None
        self._quitting_for_update = False

        self._build_ui()
        self.apply_theme(self.mode)
        self._load_settings()
        self.refresh_voices()
        self._refresh_chips()
        self._restore_geometry()
        self.log(f"🎙 {APP_NAME} {APP_VERSION} sẵn sàng. Dữ liệu lưu tại: {APP_DIR}")
        self.update_timer = QTimer(self)
        self.update_timer.timeout.connect(lambda: self.check_updates(silent=True))
        if self.config.get("auto_update", True) and not os.getenv("TTS_SELFCHECK") and not os.getenv("TTS_NO_UPDATE_CHECK"):
            QTimer.singleShot(6000, lambda: self.check_updates(silent=True))
            self.update_timer.start(UPDATE_EVERY_MS)
        if not self.voice_store.list():
            self.log("💡 Bắt đầu: vào ⚙ Cài đặt API nhập key → 🧬 Clone giọng (hoặc ☁ lấy giọng từ tài khoản).")

    # ================================================================ layout
    def _build_ui(self):
        central = QWidget()
        central.setObjectName("Central")
        root = QVBoxLayout(central)
        root.setContentsMargins(14, 12, 14, 8)
        root.setSpacing(12)
        root.addWidget(self._header())

        body = QHBoxLayout()
        body.setSpacing(12)
        body.addWidget(self._sidebar())

        self.stack = QStackedWidget()
        for builder in (self._page_clone, self._page_tts, self._page_batch, self._page_video,
                        self._page_settings, self._page_update, self._page_guide):
            self.stack.addWidget(builder())

        splitter = QSplitter(Qt.Orientation.Vertical)
        splitter.setChildrenCollapsible(False)
        splitter.addWidget(self.stack)
        splitter.addWidget(self._log_panel())
        splitter.setStretchFactor(0, 1)
        splitter.setStretchFactor(1, 0)
        splitter.setSizes([740, 110])
        self.splitter = splitter
        body.addWidget(splitter, 1)
        root.addLayout(body, 1)
        self.setCentralWidget(central)

        self.status_label = QLabel("Sẵn sàng")
        self.statusBar().addWidget(self.status_label, 1)
        self.statusBar().addPermanentWidget(QLabel(f"📁 {APP_DIR}"))
        self.go(PAGE_CLONE if not self.voice_store.list() else PAGE_TTS)

    def _header(self):
        head = QFrame()
        head.setObjectName("Header")
        lay = QHBoxLayout(head)
        lay.setContentsMargins(18, 10, 16, 10)
        lay.setSpacing(12)
        logo = QLabel()
        logo.setObjectName("Logo")
        pm = QPixmap(resource_path("assets/app.png"))
        if pm.isNull():
            logo.setText("🎙")
        else:
            logo.setPixmap(pm.scaled(48, 48, Qt.AspectRatioMode.KeepAspectRatio,
                                     Qt.TransformationMode.SmoothTransformation))
        lay.addWidget(logo)
        tv = QVBoxLayout()
        tv.setSpacing(0)
        t = QLabel(APP_NAME)
        t.setObjectName("AppTitle")
        s = QLabel("Clone giọng  •  Đọc văn bản  •  MP3 / MP4 hàng loạt từ Excel")
        s.setObjectName("AppSub")
        tv.addWidget(t)
        tv.addWidget(s)
        lay.addLayout(tv)
        # Chips live in a box that never forces the window wider; chips that
        # do not fit are hidden in _apply_responsive().
        self.chip_box = QWidget()
        self.chip_box.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Preferred)
        cb = QHBoxLayout(self.chip_box)
        cb.setContentsMargins(0, 0, 0, 0)
        cb.setSpacing(8)
        cb.addStretch()
        self.chip_ffmpeg = label("", chip="true")
        self.chip_inworld = label("", chip="true")
        self.chip_minimax = label("", chip="true")
        self.chip_voices = label("", chip="true")
        for c in (self.chip_voices, self.chip_inworld, self.chip_minimax, self.chip_ffmpeg):
            cb.addWidget(c)
        lay.addWidget(self.chip_box, 1)
        self.update_chip = QPushButton("⬆  Có bản mới")
        self.update_chip.setObjectName("UpdateChip")
        self.update_chip.setCursor(Qt.CursorShape.PointingHandCursor)
        self.update_chip.setToolTip("Bấm để xem và cài bản mới")
        self.update_chip.clicked.connect(lambda: self.go(PAGE_UPDATE))
        self.update_chip.hide()
        lay.addWidget(self.update_chip)
        self.theme_btn = QPushButton()
        self.theme_btn.setObjectName("ThemeBtn")
        self.theme_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.theme_btn.clicked.connect(lambda: self.apply_theme("light" if self.mode == "dark" else "dark", save=True))
        lay.addWidget(self.theme_btn)
        return head

    def _sidebar(self):
        side = QFrame()
        side.setObjectName("Sidebar")
        side.setFixedWidth(212)
        lay = QVBoxLayout(side)
        lay.setContentsMargins(10, 14, 10, 12)
        lay.setSpacing(4)
        cap = QLabel("  CHỨC NĂNG")
        cap.setObjectName("SideCaption")
        lay.addWidget(cap)
        self.nav_buttons: list[QPushButton] = []
        for i, (icon, text, accent) in enumerate(NAV):
            if i == PAGE_SETTINGS:
                lay.addSpacing(10)
                c2 = QLabel("  HỆ THỐNG")
                c2.setObjectName("SideCaption")
                lay.addWidget(c2)
            b = QPushButton(f"{icon}   {text}")
            b.setProperty("nav", True)
            b.setProperty("accent", accent)
            b.setCheckable(True)
            b.setCursor(Qt.CursorShape.PointingHandCursor)
            b.clicked.connect(lambda _=False, idx=i: self.go(idx))
            lay.addWidget(b)
            self.nav_buttons.append(b)
        lay.addStretch()
        tip = label("💡 Mẹo: giữ chuột trên nút để xem giải thích.", role="hint", wrap=True)
        lay.addWidget(tip)
        v = QLabel(f"  Phiên bản {APP_VERSION}")
        v.setObjectName("Version")
        lay.addWidget(v)
        return side

    def go(self, idx: int):
        self.stack.setCurrentIndex(idx)
        if idx == PAGE_VIDEO:
            QTimer.singleShot(0, self._update_video_preview)
        for i, b in enumerate(self.nav_buttons):
            b.setChecked(i == idx)

    def _page(self, icon: str, title: str, desc: str, accent: str):
        page = QWidget()
        page.setObjectName("Page")
        outer = QVBoxLayout(page)
        outer.setContentsMargins(0, 0, 0, 0)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)
        inner = QWidget()
        inner.setObjectName("Page")
        lay = QVBoxLayout(inner)
        lay.setContentsMargins(4, 2, 8, 6)
        lay.setSpacing(12)
        head = QHBoxLayout()
        head.setSpacing(12)
        badge = QLabel(icon)
        badge.setAlignment(Qt.AlignmentFlag.AlignCenter)
        badge.setFixedSize(46, 46)
        main, light, dark = ACCENTS[accent]
        badge.setStyleSheet(f"font-size:20pt; border-radius:12px; background: qlineargradient(x1:0,y1:0,x2:1,y2:1,"
                            f" stop:0 {light}, stop:1 {dark});")
        head.addWidget(badge)
        tv = QVBoxLayout()
        tv.setSpacing(0)
        t = QLabel(title)
        t.setObjectName("PageTitle")
        d = QLabel(desc)
        d.setObjectName("PageDesc")
        d.setWordWrap(True)
        tv.addWidget(t)
        tv.addWidget(d)
        head.addLayout(tv, 1)
        lay.addLayout(head)
        scroll.setWidget(inner)
        outer.addWidget(scroll)
        return page, lay, head

    # ---------------------------------------------------------------- clone
    def _page_clone(self):
        page, lay, _ = self._page("🧬", "Clone giọng nói",
                                  "Tải lên một đoạn giọng mẫu → nhận Voice ID dùng lại mãi mãi cho mọi lần đọc.",
                                  "violet")
        row = ResponsiveRow(1000)
        self.responsive.append(row)

        form = Card("Tạo giọng mới", "violet")
        form.add(label("① Chọn nhà cung cấp", role="field"))
        self.clone_provider = ChoiceRow([("Inworld", "🌊  Inworld", "teal"), ("MiniMax", "🌸  MiniMax", "pink")])
        self.clone_provider.changed.connect(self._clone_provider_changed)
        form.add(self.clone_provider)
        self.clone_rules = label("", role="hint", wrap=True)
        form.add(self.clone_rules)

        form.add(label("② File giọng mẫu", role="field"))
        sr = QHBoxLayout()
        self.clone_sample = QLineEdit()
        self.clone_sample.setPlaceholderText("Chọn file WAV / MP3 / M4A…")
        self.clone_sample.editingFinished.connect(self._validate_sample)
        sr.addWidget(self.clone_sample, 1)
        sr.addWidget(button("📂 Chọn", tint="violet", slot=self._browse_clone_sample, tip="Chọn file giọng mẫu"))
        sr.addWidget(button("▶", tint="cyan", tip="Nghe file mẫu bằng trình phát mặc định",
                            slot=lambda: self.open_path(self.clone_sample.text())))
        form.add(sr)
        self.clone_sample_info = label("Mẫu tốt nhất: 1 người nói, rõ, không nhạc nền, không vang.",
                                       role="hint", wrap=True)
        form.add(self.clone_sample_info)

        form.add(label("③ Thông tin giọng", role="field"))
        g = QGridLayout()
        g.setHorizontalSpacing(10)
        g.setVerticalSpacing(8)
        self.clone_name = QLineEdit()
        self.clone_name.setPlaceholderText("Ví dụ: Giọng nam MC")
        self.clone_language = QComboBox()
        g.addWidget(label("Tên giọng", role="muted"), 0, 0)
        g.addWidget(self.clone_name, 0, 1)
        g.addWidget(label("Ngôn ngữ", role="muted"), 1, 0)
        g.addWidget(self.clone_language, 1, 1)
        form.add(g)
        self.clone_denoise = QCheckBox("🔇  Lọc tạp âm nền trong file mẫu")
        form.add(self.clone_denoise)
        self.clone_consent = QCheckBox("Tôi là chủ giọng nói này hoặc đã được\nngười nói đồng ý cho phép clone.")
        form.add(self.clone_consent)
        self.clone_button = button("🧬   Bắt đầu Clone giọng", variant="violet", size="big", slot=self.start_clone)
        form.add(self.clone_button)
        self.clone_busy = QProgressBar()
        self.clone_busy.setRange(0, 0)
        self.clone_busy.setTextVisible(False)
        self.clone_busy.setFixedHeight(8)
        self.clone_busy.hide()
        form.add(self.clone_busy)
        form.body.addStretch()
        row.add(form, 0, 410)

        lib = Card("Thư viện giọng của bạn", "blue")
        self.voice_count_badge = label("0", badge="blue")
        lib.head.insertWidget(2, self.voice_count_badge)
        self.voice_table = QTableWidget(0, 5)
        self.voice_table.setHorizontalHeaderLabels(["Tên giọng", "Nhà cung cấp", "Voice ID", "Ngôn ngữ", "Ngày tạo"])
        hh = self.voice_table.horizontalHeader()
        hh.setSectionResizeMode(0, QHeaderView.ResizeMode.Interactive)
        hh.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        hh.setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        hh.setSectionResizeMode(3, QHeaderView.ResizeMode.ResizeToContents)
        hh.setSectionResizeMode(4, QHeaderView.ResizeMode.ResizeToContents)
        self.voice_table.setColumnWidth(0, 160)
        self.voice_table.verticalHeader().setVisible(False)
        self.voice_table.setAlternatingRowColors(True)
        self.voice_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.voice_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.voice_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.voice_table.setMinimumHeight(260)
        self.voice_table.doubleClicked.connect(lambda _: self.use_selected_voice())
        lib.add(self.voice_table)
        self.voice_empty = label("📭 Chưa có giọng nào. Clone giọng mới ở bên trái, hoặc bấm ☁ Lấy từ tài khoản.",
                                 role="hint", wrap=True)
        lib.add(self.voice_empty)
        r1 = QHBoxLayout()
        r1.addWidget(button("🗣  Dùng giọng này", variant="blue", slot=self.use_selected_voice,
                            tip="Chọn giọng này và chuyển sang trang Đọc văn bản (hoặc nháy đúp vào dòng)"))
        r1.addWidget(button("▶  Nghe thử", tint="cyan", slot=self.preview_selected_voice,
                            tip="Đọc một câu mẫu bằng giọng đang chọn"))
        r1.addWidget(button("✏  Đổi tên", tint="amber", slot=self.rename_selected_voice))
        r1.addWidget(button("📋  Sao chép ID", tint="violet", slot=self.copy_selected_voice_id))
        r1.addStretch()
        lib.add(r1)
        r2 = QHBoxLayout()
        r2.addWidget(button("➕  Thêm Voice ID", tint="blue", slot=self.add_voice_manual,
                            tip="Thêm Voice ID đã có sẵn (tạo trên web hoặc máy khác)"))
        r2.addWidget(button("☁  Lấy từ Inworld", tint="teal", slot=lambda: self.import_voices("Inworld"),
                            tip="Tải danh sách giọng trên tài khoản Inworld của bạn"))
        r2.addWidget(button("☁  Lấy từ MiniMax", tint="pink", slot=lambda: self.import_voices("MiniMax"),
                            tip="Tải danh sách giọng trên tài khoản MiniMax của bạn"))
        r2.addStretch()
        r2.addWidget(button("🗑  Xóa", tint="red", slot=self.remove_selected_voice,
                            tip="Chỉ xóa khỏi danh sách trong app, giọng trên tài khoản vẫn còn"))
        lib.add(r2)
        row.add(lib, 1)
        lay.addWidget(row, 1)
        self.clone_provider.set_value(self.config.get("last_provider", "Inworld"))
        self._clone_provider_changed(self.clone_provider.value())
        return page

    # ---------------------------------------------------------------- tts
    def _page_tts(self):
        page, lay, _ = self._page("🗣", "Đọc văn bản",
                                  "Nhập hoặc dán nội dung, chọn giọng rồi bấm Tạo giọng đọc. Văn bản dài được tự chia đoạn và ghép lại.",
                                  "blue")
        cols = ResponsiveRow(860)
        self.responsive.append(cols)

        # ---- left: text
        text_card = Card("Nội dung cần đọc", "cyan")
        tools = QHBoxLayout()
        tools.addWidget(button("📋 Dán", size="small", tint="cyan", slot=self._paste_text))
        tools.addWidget(button("📄 Mở .txt", size="small", tint="teal", slot=self._open_txt))
        tools.addWidget(button("🧹 Xóa hết", size="small", tint="red", slot=lambda: self.tts_text.clear()))
        text_card.head.addLayout(tools)
        self.tts_text = QTextEdit()
        self.tts_text.setAcceptRichText(False)
        self.tts_text.setPlaceholderText("✍  Nhập hoặc dán văn bản vào đây…\n\nVí dụ: Xin chào các bạn, chào mừng đến với kênh của mình!")
        self.tts_text.setMinimumHeight(120)
        f = self.tts_text.font()
        f.setPointSize(11)
        self.tts_text.setFont(f)
        self.tts_text.textChanged.connect(self._update_counter)
        text_card.add(self.tts_text)
        self.tts_counter = label("", role="hint")
        text_card.add(self.tts_counter)
        ar = QHBoxLayout()
        ar.setSpacing(8)
        self.tts_generate = button("🔊   Tạo giọng đọc", variant="blue", size="big", slot=self.start_tts, min_w=220)
        self.tts_stop = button("⏹ Dừng", variant="red", size="big", slot=self.stop_tts)
        self.tts_stop.setEnabled(False)
        self.tts_open_file = button("▶ Mở file", tint="green", slot=self._open_last_output,
                                    tip="Mở file vừa tạo bằng trình phát mặc định")
        self.tts_open_file.setEnabled(False)
        ar.addWidget(self.tts_generate)
        ar.addWidget(self.tts_stop)
        ar.addStretch()
        ar.addWidget(self.tts_open_file)
        ar.addWidget(button("📂 Thư mục", tint="amber", slot=lambda: self.open_folder(self.tts_output_dir.text())))
        text_card.add(ar)
        self.tts_progress = QProgressBar()
        self.tts_progress.setValue(0)
        text_card.add(self.tts_progress)
        self.tts_result = label("", role="hint", wrap=True)
        self.tts_result.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        text_card.add(self.tts_result)
        cols.add(text_card, 1)

        # ---- right: settings + actions
        right = QVBoxLayout()
        right.setSpacing(12)
        top = Card("Giọng & thông số", "blue")
        g = QGridLayout()
        g.setHorizontalSpacing(10)
        g.setVerticalSpacing(8)
        self.tts_voice = QComboBox()
        self.tts_voice.currentIndexChanged.connect(lambda _: (self._voice_changed(
            self.tts_voice, self.tts_model, self.tts_speed, self.tts_voice_info), self._update_counter()))
        self.tts_model = QComboBox()
        self.tts_model.setEditable(True)
        self.tts_speed = QDoubleSpinBox()
        self.tts_speed.setDecimals(2)
        self.tts_speed.setSingleStep(0.05)
        self.tts_speed.setSuffix(" ×")
        self.tts_speed.setValue(1.0)
        self.tts_voice_info = label("", role="hint", wrap=True)
        g.addWidget(label("🎤 Giọng", role="field"), 0, 0)
        g.addWidget(self.tts_voice, 0, 1)
        g.addWidget(label("🧠 Model", role="field"), 1, 0)
        g.addWidget(self.tts_model, 1, 1)
        g.addWidget(label("⚡ Tốc độ", role="field"), 2, 0)
        g.addWidget(self.tts_speed, 2, 1)
        g.addWidget(self.tts_voice_info, 3, 0, 1, 2)
        g.setColumnStretch(1, 1)
        top.add(g)
        right.addWidget(top)

        out = Card("Xuất file", "green")
        og = QGridLayout()
        og.setHorizontalSpacing(10)
        og.setVerticalSpacing(8)
        self.tts_output_dir = QLineEdit()
        self.tts_filename = QLineEdit()
        self.tts_filename.setPlaceholderText("Trống = tự đặt theo giờ")
        self.tts_format = ChoiceRow(FORMAT_OPTIONS)
        self.tts_format.changed.connect(self._format_changed)
        og.addWidget(label("📁 Thư mục", role="field"), 0, 0)
        dr = QHBoxLayout()
        dr.addWidget(self.tts_output_dir, 1)
        dr.addWidget(button("📂", tint="green", tip="Chọn thư mục lưu",
                            slot=lambda: self._pick_folder(self.tts_output_dir)))
        og.addLayout(dr, 0, 1)
        og.addWidget(label("🏷 Tên file", role="field"), 1, 0)
        og.addWidget(self.tts_filename, 1, 1)
        og.addWidget(label("💾 Định dạng", role="field"), 2, 0)
        og.addWidget(self.tts_format, 2, 1)
        og.addWidget(button("🎬 Ảnh nền, khung hình video…", size="small", tint="pink",
                            slot=lambda: self.go(PAGE_VIDEO)), 3, 1)
        og.setColumnStretch(1, 1)
        out.add(og)
        right.addWidget(out)

        right.addStretch()
        rw = QWidget()
        rw.setLayout(right)
        right.setContentsMargins(0, 0, 0, 0)
        cols.add(rw, 0, 430)
        lay.addWidget(cols, 1)
        self._update_counter()
        return page

    # ---------------------------------------------------------------- batch
    def _page_batch(self):
        page, lay, _ = self._page("📊", "Hàng loạt từ Excel",
                                  "Mỗi dòng Excel → 1 file giọng đọc (MP3/MP4). Có thể gộp tất cả thành 1 file.",
                                  "green")
        row = ResponsiveRow(940)
        self.responsive.append(row)

        src = Card("Nguồn Excel", "green")
        sg = QGridLayout()
        sg.setHorizontalSpacing(10)
        sg.setVerticalSpacing(8)
        self.batch_file = QLineEdit()
        self.batch_file.setReadOnly(True)
        self.batch_file.setPlaceholderText("Chưa chọn file .xlsx")
        self.batch_sheet = QComboBox()
        self.batch_sheet.currentTextChanged.connect(self._batch_sheet_changed)
        self.batch_text_col = QComboBox()
        self.batch_name_col = QComboBox()
        self.batch_text_col.currentIndexChanged.connect(lambda _: self._batch_load_tasks())
        self.batch_name_col.currentIndexChanged.connect(lambda _: self._batch_load_tasks())
        sg.addWidget(label("📄 File", role="field"), 0, 0)
        sg.addWidget(self.batch_file, 0, 1)
        sg.addWidget(button("📂 Chọn", tint="green", slot=self.select_batch_file), 0, 2)
        sg.addWidget(label("📑 Sheet", role="field"), 1, 0)
        sg.addWidget(self.batch_sheet, 1, 1, 1, 2)
        sg.addWidget(label("📝 Cột nội dung", role="field"), 2, 0)
        sg.addWidget(self.batch_text_col, 2, 1, 1, 2)
        sg.addWidget(label("🏷 Cột tên file", role="field"), 3, 0)
        sg.addWidget(self.batch_name_col, 3, 1, 1, 2)
        sg.setColumnStretch(1, 1)
        src.add(sg)
        br = QHBoxLayout()
        br.addWidget(button("📥 Tạo file Excel mẫu", size="small", tint="teal", slot=self.create_sample_excel))
        br.addWidget(button("🔄 Tải lại", size="small", tint="blue", slot=self._reload_batch_file))
        br.addWidget(button("📗 Mở Excel", size="small", tint="green",
                            slot=lambda: self.open_path(self.batch_path)))
        br.addStretch()
        src.add(br)
        src.body.addStretch()
        row.add(src, 1)

        cfg = Card("Giọng & xuất file", "orange")
        cg = QGridLayout()
        cg.setHorizontalSpacing(10)
        cg.setVerticalSpacing(8)
        self.batch_voice = QComboBox()
        self.batch_model = QComboBox()
        self.batch_model.setEditable(True)
        self.batch_speed = QDoubleSpinBox()
        self.batch_speed.setDecimals(2)
        self.batch_speed.setSingleStep(0.05)
        self.batch_speed.setSuffix(" ×")
        self.batch_voice_info = label("", role="hint")
        self.batch_voice.currentIndexChanged.connect(lambda _: self._voice_changed(
            self.batch_voice, self.batch_model, self.batch_speed, self.batch_voice_info))
        self.batch_output_dir = QLineEdit()
        self.batch_format = ChoiceRow(FORMAT_OPTIONS)
        self.batch_format.changed.connect(self._format_changed)
        self.batch_threads = QSpinBox()
        self.batch_threads.setRange(1, 6)
        self.batch_threads.setSuffix(" luồng")
        self.batch_threads.setToolTip("Số dòng xử lý cùng lúc. 2–3 là an toàn; quá cao dễ bị giới hạn tốc độ API.")
        self.batch_skip = QCheckBox("Bỏ qua dòng đã có file")
        self.batch_skip.setToolTip("Chạy tiếp từ chỗ dừng: dòng nào đã có file đầu ra sẽ không tạo lại (không tốn phí).")
        self.batch_merge = QCheckBox("Gộp tất cả thành 1 file (thêm file _GOP)")
        self.batch_merge.setToolTip("Ngoài từng file riêng, tạo thêm 1 file gộp theo thứ tự dòng.")
        cg.addWidget(label("🎤 Giọng", role="field"), 0, 0)
        cg.addWidget(self.batch_voice, 0, 1, 1, 3)
        cg.addWidget(label("🧠 Model", role="field"), 1, 0)
        cg.addWidget(self.batch_model, 1, 1)
        cg.addWidget(label("⚡ Tốc độ", role="field"), 1, 2)
        cg.addWidget(self.batch_speed, 1, 3)
        cg.addWidget(label("📁 Thư mục", role="field"), 2, 0)
        cg.addWidget(self.batch_output_dir, 2, 1, 1, 2)
        cg.addWidget(button("📂", tint="orange", slot=lambda: self._pick_folder(self.batch_output_dir),
                            tip="Chọn thư mục lưu"), 2, 3)
        cg.addWidget(label("💾 Định dạng", role="field"), 3, 0)
        cg.addWidget(self.batch_format, 3, 1, 1, 3)
        opts = QHBoxLayout()
        opts.addWidget(self.batch_threads)
        opts.addSpacing(10)
        opts.addWidget(self.batch_skip)
        opts.addStretch()
        cg.addWidget(label("🛠 Tùy chọn", role="field"), 4, 0)
        cg.addLayout(opts, 4, 1, 1, 3)
        cg.addWidget(self.batch_merge, 5, 1, 1, 3)
        cg.setColumnStretch(1, 1)
        cfg.add(cg)
        cfg.body.addStretch()
        row.add(cfg, 1)
        lay.addWidget(row)

        tiles = QHBoxLayout()
        tiles.setSpacing(10)
        self.tile_total = StatTile("Tổng dòng", "blue", "📋")
        self.tile_ok = StatTile("Thành công", "green", "✅")
        self.tile_err = StatTile("Lỗi", "red", "❌")
        self.tile_skip = StatTile("Bỏ qua", "amber", "⏭")
        self.tile_chars = StatTile("Tổng ký tự", "violet", "🔤")
        for t in (self.tile_total, self.tile_ok, self.tile_err, self.tile_skip, self.tile_chars):
            tiles.addWidget(t)
        lay.addLayout(tiles)

        run = QHBoxLayout()
        run.setSpacing(10)
        self.batch_run = button("🚀   Chạy hàng loạt", variant="green", size="big", slot=self.start_batch, min_w=190)
        self.batch_stop = button("⏹  Dừng", variant="red", size="big", slot=self.stop_batch)
        self.batch_stop.setEnabled(False)
        self.batch_retry = button("🔁  Chạy lại dòng lỗi", tint="orange", slot=lambda: self.start_batch(only_failed=True),
                                  tip="Chỉ chạy lại các dòng bị lỗi hoặc đã dừng")
        self.batch_retry.setEnabled(False)
        run.addWidget(self.batch_run)
        run.addWidget(self.batch_stop)
        run.addWidget(self.batch_retry)
        run.addStretch()
        run.addWidget(button("📂  Thư mục", tint="amber", slot=lambda: self.open_folder(self.batch_output_dir.text()),
                             tip="Mở thư mục chứa file đã tạo"))
        lay.addLayout(run)
        self.batch_progress = QProgressBar()
        self.batch_progress.setProperty("accent", "green")
        self.batch_progress.setValue(0)
        lay.addWidget(self.batch_progress)
        self.batch_status = label("Chưa chạy.", role="hint")
        lay.addWidget(self.batch_status)

        self.batch_table = QTableWidget(0, 5)
        self.batch_table.setHorizontalHeaderLabels(["Dòng", "Tên file", "Nội dung", "Ký tự", "Trạng thái"])
        bh = self.batch_table.horizontalHeader()
        bh.setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        bh.setSectionResizeMode(1, QHeaderView.ResizeMode.Interactive)
        bh.setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        bh.setSectionResizeMode(3, QHeaderView.ResizeMode.ResizeToContents)
        bh.setSectionResizeMode(4, QHeaderView.ResizeMode.Interactive)
        self.batch_table.setColumnWidth(1, 150)
        self.batch_table.setColumnWidth(4, 260)
        self.batch_table.verticalHeader().setVisible(False)
        self.batch_table.setAlternatingRowColors(True)
        self.batch_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.batch_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.batch_table.setMinimumHeight(240)
        self.batch_table.setWordWrap(False)
        lay.addWidget(self.batch_table, 1)
        return page

    # ---------------------------------------------------------------- video
    def _page_video(self):
        page, lay, _ = self._page("🎬", "Video MP4",
                                  "Thiết lập khung hình cho video: ảnh nền tĩnh (hoặc màu nền) + giọng đọc. "
                                  "Áp dụng khi chọn định dạng MP4 hoặc Cả hai.",
                                  "pink")
        row = ResponsiveRow(900)
        self.responsive.append(row)
        left = QVBoxLayout()
        left.setSpacing(12)

        img = Card("Ảnh nền", "pink")
        ir = QHBoxLayout()
        self.video_image = QLineEdit()
        self.video_image.setPlaceholderText("Không dùng ảnh (chỉ màu nền)")
        self.video_image.setReadOnly(True)
        ir.addWidget(self.video_image, 1)
        ir.addWidget(button("🖼 Chọn ảnh", tint="pink", slot=self._pick_video_image))
        ir.addWidget(button("✖", tint="red", tip="Bỏ ảnh nền", slot=lambda: self._set_video_image("")))
        img.add(ir)
        img.add(label("Ảnh được co giãn giữ đúng tỉ lệ, phần thừa tô bằng màu nền. Hỗ trợ JPG, PNG, WEBP, BMP.",
                      role="hint", wrap=True))
        left.addWidget(img)

        frame = Card("Khung hình & màu nền", "violet")
        fg = QGridLayout()
        fg.setHorizontalSpacing(10)
        fg.setVerticalSpacing(8)
        self.video_size = QComboBox()
        self.video_size.addItems(list(VIDEO_SIZES))
        self.video_size.currentTextChanged.connect(self._video_size_changed)
        self.video_color_btn = button("", tint="violet", slot=self._pick_video_color)
        self.video_color = "#101828"
        fg.addWidget(label("📐 Kích thước", role="field"), 0, 0)
        fg.addWidget(self.video_size, 0, 1)
        fg.addWidget(label("🎨 Màu nền", role="field"), 1, 0)
        fg.addWidget(self.video_color_btn, 1, 1)
        fg.setColumnStretch(1, 1)
        frame.add(fg)
        left.addWidget(frame)

        ff = Card("Bộ xử lý video (ffmpeg)", "cyan")
        self.ffmpeg_label = label("", role="hint", wrap=True)
        ff.add(self.ffmpeg_label)
        fr = QHBoxLayout()
        fr.addWidget(button("🔍 Kiểm tra lại", tint="cyan", slot=self._refresh_chips))
        self.video_test_btn = button("🎞 Tạo video thử 3 giây", tint="green", slot=self.make_test_video)
        fr.addWidget(self.video_test_btn)
        fr.addStretch()
        ff.add(fr)
        left.addWidget(ff)
        left.addStretch()
        lw = QWidget()
        lw.setLayout(left)
        left.setContentsMargins(0, 0, 0, 0)
        row.add(lw, 0, 430)

        prev = Card("Xem trước khung hình", "blue")
        self.video_preview = QLabel()
        self.video_preview.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.video_preview.setMinimumSize(420, 300)
        self.video_preview.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        prev.add(self.video_preview)
        self.video_preview_info = label("", role="hint")
        self.video_preview_info.setAlignment(Qt.AlignmentFlag.AlignCenter)
        prev.add(self.video_preview_info)
        row.add(prev, 1)
        lay.addWidget(row, 1)
        return page

    # ---------------------------------------------------------------- settings
    def _key_row(self, edit: QLineEdit):
        edit.setEchoMode(QLineEdit.EchoMode.Password)
        r = QHBoxLayout()
        r.addWidget(edit, 1)
        eye = button("👁", size="small", tip="Hiện / ẩn key")
        eye.setCheckable(True)
        eye.toggled.connect(lambda on: edit.setEchoMode(QLineEdit.EchoMode.Normal if on else QLineEdit.EchoMode.Password))
        r.addWidget(eye)
        return r

    def _page_settings(self):
        page, lay, _ = self._page("⚙", "Cài đặt API",
                                  "Nhập API key của nhà cung cấp. Key được cất trong Windows Credential Manager, không ghi ra file.",
                                  "amber")
        row = ResponsiveRow(820)
        self.responsive.append(row)

        iw = Card("🌊  Inworld", "teal", "Clone 5–30 giây mẫu • 200+ ngôn ngữ • tốc độ 0.5–1.5×")
        g = QGridLayout()
        g.setHorizontalSpacing(10)
        g.setVerticalSpacing(8)
        self.iw_key = QLineEdit()
        self.iw_key.setPlaceholderText("Dán Base64 API key (Basic) từ Inworld Portal")
        self.iw_base = QLineEdit()
        g.addWidget(label("🔑 API key", role="field"), 0, 0)
        g.addLayout(self._key_row(self.iw_key), 0, 1)
        g.addWidget(label("🌐 Base URL", role="field"), 1, 0)
        g.addWidget(self.iw_base, 1, 1)
        g.setColumnStretch(1, 1)
        iw.add(g)
        r = QHBoxLayout()
        self.iw_test = button("🔌 Kiểm tra kết nối", tint="teal", slot=lambda: self.test_connection("Inworld"))
        r.addWidget(self.iw_test)
        r.addWidget(button("🔗 Lấy API key", tint="blue",
                           slot=lambda: QDesktopServices.openUrl(QUrl("https://platform.inworld.ai/"))))
        r.addStretch()
        iw.add(r)
        self.iw_test_result = label("", role="hint", wrap=True)
        iw.add(self.iw_test_result)
        iw.body.addStretch()
        row.add(iw, 1)

        mm = Card("🌸  MiniMax", "pink", "Clone 10 giây – 5 phút mẫu • tốc độ 0.5–2× • giọng clone bị xóa nếu 7 ngày không dùng")
        g2 = QGridLayout()
        g2.setHorizontalSpacing(10)
        g2.setVerticalSpacing(8)
        self.mm_key = QLineEdit()
        self.mm_key.setPlaceholderText("Dán API key (Bearer) từ MiniMax")
        self.mm_group = QLineEdit()
        self.mm_group.setPlaceholderText("Không bắt buộc")
        self.mm_base = QComboBox()
        self.mm_base.setEditable(True)
        self.mm_base.addItems(["https://api.minimax.io", "https://api-uw.minimax.io", "https://api.minimaxi.com"])
        self.mm_base.setToolTip("api.minimax.io = quốc tế • api-uw = máy chủ Mỹ • api.minimaxi.com = Trung Quốc")
        g2.addWidget(label("🔑 API key", role="field"), 0, 0)
        g2.addLayout(self._key_row(self.mm_key), 0, 1)
        g2.addWidget(label("👥 Group ID", role="field"), 1, 0)
        g2.addWidget(self.mm_group, 1, 1)
        g2.addWidget(label("🌐 Base URL", role="field"), 2, 0)
        g2.addWidget(self.mm_base, 2, 1)
        g2.setColumnStretch(1, 1)
        mm.add(g2)
        r2 = QHBoxLayout()
        self.mm_test = button("🔌 Kiểm tra kết nối", tint="pink", slot=lambda: self.test_connection("MiniMax"))
        r2.addWidget(self.mm_test)
        r2.addWidget(button("🔗 Lấy API key", tint="blue", slot=lambda: QDesktopServices.openUrl(
            QUrl("https://platform.minimax.io/user-center/basic-information/interface-key"))))
        r2.addStretch()
        mm.add(r2)
        self.mm_test_result = label("", role="hint", wrap=True)
        mm.add(self.mm_test_result)
        mm.body.addStretch()
        row.add(mm, 1)
        lay.addWidget(row)

        bottom = Card("Lưu & dữ liệu", "amber")
        br = QHBoxLayout()
        br.addWidget(button("💾   Lưu cài đặt", variant="amber", size="big", slot=self.save_settings, min_w=200))
        br.addWidget(button("📁  Mở thư mục dữ liệu", tint="blue", slot=lambda: self.open_folder(str(APP_DIR))))
        br.addWidget(button("📝  Mở file log", tint="violet", slot=lambda: self.open_path(str(LOG_PATH))))
        br.addStretch()
        bottom.add(br)
        bottom.add(label("🔒 API key lưu trong Windows Credential Manager (keyring). Danh sách giọng và cài đặt lưu ở "
                         f"{APP_DIR}. Có thể đặt biến môi trường INWORLD_API_KEY / MINIMAX_API_KEY thay cho việc nhập tay.",
                         role="hint", wrap=True))
        lay.addWidget(bottom)
        lay.addStretch()
        return page

    # ---------------------------------------------------------------- guide
    def _page_update(self):
        page, lay, _ = self._page("🔄", "Cập nhật phần mềm",
                                  "Tự kiểm tra, tải và cài bản mới từ GitHub — không cần tải lại bằng tay.", "teal")
        row = ResponsiveRow(900)
        self.responsive.append(row)
        info = Card("Phiên bản", "teal")
        g = QGridLayout()
        g.setHorizontalSpacing(12)
        g.setVerticalSpacing(8)
        cur = label(APP_VERSION, role="cardTitle", accent="teal")
        cur.setStyleSheet("font-size: 22pt; font-weight: 800;")
        self.upd_kind = label(updater.KIND_TEXT.get(updater.install_kind(), ""), role="muted", wrap=True)
        repo_txt = (f"<a href='https://github.com/{GITHUB_REPO}/releases'>github.com/{GITHUB_REPO}</a>"
                    if updater.repo_configured() else "(chưa cấu hình)")
        repo = QLabel(repo_txt)
        repo.setOpenExternalLinks(True)
        repo.setWordWrap(True)
        self.upd_last = label(self.config.get("last_update_check") or "Chưa kiểm tra", role="muted")
        for r, (t, w) in enumerate([("📦 Đang dùng", cur), ("💻 Kiểu cài đặt", self.upd_kind),
                                    ("☁ Nguồn cập nhật", repo), ("🕒 Kiểm tra lần cuối", self.upd_last)]):
            g.addWidget(label(t, role="field"), r, 0, Qt.AlignmentFlag.AlignTop)
            g.addWidget(w, r, 1)
        g.setColumnStretch(1, 1)
        info.add(g)
        self.upd_auto = QCheckBox("Tự kiểm tra bản mới khi mở ứng dụng (và 6 giờ/lần)")
        self.upd_auto.setChecked(bool(self.config.get("auto_update", True)))
        self.upd_auto.toggled.connect(lambda on: setattr(self, "config", self.config_store.save({"auto_update": on})))
        info.add(self.upd_auto)
        self.upd_check_btn = button("🔍   Kiểm tra cập nhật", variant="teal", size="big",
                                    slot=lambda: self.check_updates(silent=False))
        info.add(self.upd_check_btn)
        self.upd_install_btn = button("⬇   Tải && cài bản mới", variant="green", size="big", slot=self.install_update)
        self.upd_install_btn.setEnabled(False)
        info.add(self.upd_install_btn)
        self.upd_progress = QProgressBar()
        self.upd_progress.setProperty("accent", "green")
        self.upd_progress.hide()
        info.add(self.upd_progress)
        self.upd_status = label("", role="hint", wrap=True)
        info.add(self.upd_status)
        br = QHBoxLayout()
        br.addWidget(button("🌐 Trang phát hành", tint="blue", slot=self._open_release_page))
        br.addWidget(button("📝 Nhật ký cập nhật", tint="violet",
                            slot=lambda: self.open_path(str(APP_DIR / "updates" / "update_log.txt"))))
        br.addStretch()
        info.add(br)
        info.body.addStretch()
        row.add(info, 0, 430)

        notes = Card("Có gì mới", "violet")
        self.upd_notes = QTextBrowser()
        self.upd_notes.setOpenExternalLinks(True)
        self.upd_notes.setMinimumHeight(260)
        self.upd_notes.setMarkdown("Bấm **🔍 Kiểm tra cập nhật** để xem ghi chú của bản mới nhất.")
        notes.add(self.upd_notes)
        row.add(notes, 1)
        lay.addWidget(row, 1)
        return page

    # ================================================================ updates
    def _open_release_page(self):
        url = (self.latest_release.html_url if self.latest_release
               else f"https://github.com/{GITHUB_REPO}/releases")
        QDesktopServices.openUrl(QUrl(url))

    def _busy(self) -> bool:
        return any(w.isRunning() for w in list(self.workers) if w is not self.update_worker)

    def check_updates(self, silent: bool = False):
        if self.update_worker and self.update_worker.isRunning():
            return
        if not updater.repo_configured():
            self.upd_status.setText("⚠ Bản này chưa được cấu hình kho GitHub để cập nhật.")
            return
        self.upd_check_btn.setEnabled(False)
        self.upd_status.setText("⏳ Đang kiểm tra bản mới trên GitHub…")
        self.upd_status.setStyleSheet("")

        def job(log, cancelled, progress):
            return updater.fetch_latest()

        def ok(rel):
            self.latest_release = rel
            now = datetime.now().strftime("%d/%m/%Y %H:%M")
            self.config = self.config_store.save({"last_update_check": now})
            self.upd_last.setText(now)
            notes = rel.notes.strip() or "_(Bản này không có ghi chú)_"
            self.upd_notes.setMarkdown(f"## Phiên bản {rel.version}\n\n{notes}")
            sc = STATUS_COLORS[self.mode]
            if updater.is_newer(rel.version):
                kind = updater.install_kind()
                self.upd_status.setText(f"🎉 Có bản mới {rel.version} (bạn đang dùng {APP_VERSION}).")
                self.upd_status.setStyleSheet(f"color:{sc['ok']}; font-weight:700;")
                self.upd_install_btn.setEnabled(True)
                self.upd_install_btn.setText("🌐   Tải bản mới trên GitHub" if kind == "source"
                                             else f"⬇   Tải && cài bản {rel.version}")
                self.update_chip.setText(f"⬆  Có bản mới {rel.version}")
                self.update_chip.show()
                self.log(f"🎉 Có bản mới {rel.version} — xem trang 🔄 Cập nhật.")
                if silent and rel.version != self.config.get("skip_version") and not self._busy():
                    self._ask_update(rel)
            else:
                self.upd_status.setText(f"✅ Bạn đang dùng bản mới nhất ({APP_VERSION}).")
                self.upd_status.setStyleSheet(f"color:{sc['ok']}; font-weight:600;")
                self.upd_install_btn.setEnabled(False)
                self.update_chip.hide()
                if not silent:
                    self.log(f"✅ Đang dùng bản mới nhất ({APP_VERSION}).")

        def bad(err):
            self.upd_status.setText("❌ " + err)
            self.upd_status.setStyleSheet(f"color:{STATUS_COLORS[self.mode]['error']};")
            if not silent:
                self.log("❌ Kiểm tra cập nhật: " + err)

        def fin():
            self.upd_check_btn.setEnabled(True)

        self.update_worker = self._run(job, ok, bad, finished=fin)

    def _ask_update(self, rel):
        box = QMessageBox(QMessageBox.Icon.Information, "Có bản cập nhật mới",
                          f"🎉 Đã có TTS Clone Studio {rel.version} (bạn đang dùng {APP_VERSION}).\n\n"
                          "Cập nhật ngay? Ứng dụng sẽ tự tải, cài và mở lại.", parent=self)
        now_btn = box.addButton("⬇ Cập nhật ngay", QMessageBox.ButtonRole.AcceptRole)
        box.addButton("Để sau", QMessageBox.ButtonRole.RejectRole)
        skip_btn = box.addButton("Bỏ qua bản này", QMessageBox.ButtonRole.DestructiveRole)
        box.exec()
        if box.clickedButton() is now_btn:
            self.go(PAGE_UPDATE)
            self.install_update(confirm=False)
        elif box.clickedButton() is skip_btn:
            self.config = self.config_store.save({"skip_version": rel.version})

    def install_update(self, confirm: bool = True):
        rel = self.latest_release
        if not rel or not updater.is_newer(rel.version):
            return
        kind = updater.install_kind()
        if kind == "source":
            QMessageBox.information(self, "Chạy từ mã nguồn",
                                    "Bạn đang chạy từ mã nguồn Python nên không tự cài được.\n"
                                    "Trang GitHub sẽ mở ra: tải bản mới (Source code zip) và giải nén đè lên thư mục cũ.")
            self._open_release_page()
            return
        if self._busy():
            QMessageBox.warning(self, "Đang có tác vụ chạy",
                                "Hãy chờ tác vụ đang chạy (đọc văn bản / hàng loạt / clone) xong rồi cập nhật.")
            return
        asset = updater.pick_asset(rel, kind)
        if not asset:
            want = updater.asset_name(rel.version, kind)
            QMessageBox.warning(self, "Thiếu file cập nhật",
                                f"Bản {rel.version} chưa có file cho máy này ({want}).\nTrang GitHub sẽ mở ra để tải tay.")
            self._open_release_page()
            return
        size_mb = (asset.get("size") or 0) / 1024 / 1024
        if confirm and QMessageBox.question(
                self, "Cập nhật",
                f"Tải bản {rel.version} ({size_mb:.0f} MB)?\n\nTải xong, ứng dụng sẽ tự đóng, cài bản mới và mở lại "
                "(khoảng 10–30 giây).") != QMessageBox.StandardButton.Yes:
            return
        dest = APP_DIR / "updates"
        self.upd_install_btn.setEnabled(False)
        self.upd_check_btn.setEnabled(False)
        self.upd_progress.setRange(0, 100)
        self.upd_progress.setValue(0)
        self.upd_progress.show()
        self.upd_status.setText(f"⬇ Đang tải {asset['name']}…")
        self.log(f"⬇ Đang tải bản cập nhật {rel.version}…")

        def job(log, cancelled, progress):
            path = updater.download_asset(rel, asset, dest,
                                          progress=lambda g, t: progress(int(g * 100 / t)) if t else None,
                                          cancelled=cancelled)
            if kind in ("mac-app", "mac-manual"):
                staging = dest / f"mac_{rel.version}"
                app_path = updater.extract_mac_zip(path, str(staging))
                return {"new": str(app_path), "staging": str(staging), "file": path}
            return {"new": path, "staging": "", "file": path}

        def ok(res):
            self.upd_progress.setValue(100)
            if kind == "mac-manual":
                subprocess.Popen(["open", "-R", res["new"]])
                QMessageBox.information(self, "Đã tải bản mới",
                                        "Ứng dụng đang chạy từ ngoài thư mục Applications nên không tự thay được.\n"
                                        "Finder đã mở bản mới — hãy kéo TTSCloneStudio vào Applications (chọn Thay thế).")
                return
            try:
                target = updater.current_target(kind)
                updater.launch_helper(kind, new_path=res["new"], target=target,
                                      log=str(dest / "update_log.txt"), staging=res["staging"])
            except Exception as exc:  # noqa: BLE001
                self.upd_status.setText(f"❌ Không khởi chạy được trình cài: {exc}")
                self.log(f"❌ Không khởi chạy được trình cài: {exc}")
                return
            self.upd_status.setText("✅ Đã tải xong — đang đóng ứng dụng để cài bản mới…")
            self.log("🔄 Đang cài bản mới, ứng dụng sẽ tự mở lại sau ít giây…")
            self._quitting_for_update = True
            QTimer.singleShot(400, QApplication.instance().quit)

        def bad(err):
            self.upd_progress.hide()
            self.upd_status.setText("❌ " + err)
            self.upd_status.setStyleSheet(f"color:{STATUS_COLORS[self.mode]['error']};")
            self.log("❌ Cập nhật lỗi: " + err)

        def fin():
            self.upd_check_btn.setEnabled(True)
            if not self._quitting_for_update and self.latest_release and updater.is_newer(self.latest_release.version):
                self.upd_install_btn.setEnabled(True)

        self.update_worker = self._run(job, ok, bad, progress=self.upd_progress.setValue, finished=fin)

    def _page_guide(self):
        page, lay, _ = self._page("📖", "Hướng dẫn nhanh", "4 bước để có file giọng đọc đầu tiên.", "cyan")
        self.guide = QTextBrowser()
        self.guide.setObjectName("Guide")
        self.guide.setOpenExternalLinks(True)
        self.guide.setMinimumHeight(430)
        lay.addWidget(self.guide, 1)
        return page

    def _guide_html(self):
        c = {k: (v[0] if self.mode == "dark" else v[2]) for k, v in ACCENTS.items()}

        def step(n, color, title, body):
            return (f"<table width='100%' cellspacing='0' cellpadding='6'><tr>"
                    f"<td width='42' valign='top'><div style='background:{color};color:white;font-weight:800;"
                    f"font-size:15pt;padding:4px 0;text-align:center'>&nbsp;{n}&nbsp;</div></td>"
                    f"<td><b style='font-size:12pt;color:{color}'>{title}</b><br>{body}</td></tr></table>")

        return (
            "<div style='font-size:10.5pt;line-height:150%'>"
            + step(1, c["amber"], "⚙ Nhập API key",
                   "Vào <b>Cài đặt API</b>, dán key của Inworld và/hoặc MiniMax, bấm <b>🔌 Kiểm tra kết nối</b> rồi <b>💾 Lưu</b>.")
            + step(2, c["violet"], "🧬 Tạo giọng",
                   "Trang <b>Clone giọng</b>: chọn nhà cung cấp, chọn file mẫu (1 người nói, rõ, không nhạc), đặt tên, "
                   "tích xác nhận quyền rồi bấm <b>Bắt đầu Clone</b>. Đã có giọng trên tài khoản? Bấm <b>☁ Lấy từ tài khoản</b>.")
            + step(3, c["blue"], "🗣 Đọc văn bản",
                   "Chọn giọng, dán nội dung, chọn <b>MP3</b>, <b>MP4</b> hoặc <b>Cả hai</b> rồi bấm <b>🔊 Tạo giọng đọc</b>. "
                   "Văn bản dài được tự chia đoạn và ghép liền mạch.")
            + step(4, c["green"], "📊 Hàng loạt từ Excel",
                   "Dòng đầu là tiêu đề, ví dụ cột <b>filename</b> và <b>text</b> (bấm <b>📥 Tạo file Excel mẫu</b>). "
                   "Chọn cột, bấm <b>🚀 Chạy hàng loạt</b>. Bật <b>⏭ Bỏ qua dòng đã có file</b> để chạy tiếp khi bị ngắt; "
                   "bật <b>🔗 Gộp</b> để có thêm 1 file tổng.")
            + f"<p style='color:{c['pink']}'><b>🎬 Video MP4:</b> chọn ảnh nền và khung hình ở trang <b>Video MP4</b>. "
              "Video = ảnh tĩnh + giọng đọc, chuẩn H.264/AAC phát được trên YouTube, Facebook, TikTok.</p>"
            + f"<p style='color:{c['teal']}'><b>🔄 Cập nhật:</b> ứng dụng tự kiểm tra bản mới khi mở. Có bản mới sẽ hiện nút "
              "<b>⬆ Có bản mới</b> trên thanh tiêu đề — bấm là tự tải, cài và mở lại. Dữ liệu, giọng và API key được giữ nguyên.</p>"
            + f"<p style='color:{c['orange']}'><b>⚠ Lưu ý:</b> giọng clone của MiniMax sẽ bị xóa nếu 7 ngày không dùng — "
              "hãy bấm <b>▶ Nghe thử</b> ngay sau khi clone. Chỉ clone giọng của bạn hoặc giọng đã được cho phép.</p>"
            + f"<p style='color:{c['cyan']}'><b>🛟 Gặp lỗi?</b> Xem ô <b>Nhật ký</b> bên dưới, hoặc gửi file "
              f"<code>{html.escape(str(LOG_PATH))}</code> cho người hỗ trợ.</p></div>"
        )

    # ---------------------------------------------------------------- log
    def _log_panel(self):
        card = QFrame()
        card.setObjectName("LogCard")
        lay = QVBoxLayout(card)
        lay.setContentsMargins(12, 8, 12, 10)
        lay.setSpacing(6)
        head = QHBoxLayout()
        head.addWidget(label("📝  Nhật ký hoạt động", role="cardTitle", accent="cyan"))
        head.addStretch()
        head.addWidget(button("📋 Sao chép", size="small", tint="blue",
                              slot=lambda: QGuiApplication.clipboard().setText(self.log_text.toPlainText())))
        head.addWidget(button("🧹 Xóa", size="small", tint="red", slot=lambda: self.log_text.clear()))
        self.log_toggle = button("▾ Thu gọn", size="small", tint="violet", slot=self._toggle_log)
        head.addWidget(self.log_toggle)
        lay.addLayout(head)
        self.log_text = QTextEdit()
        self.log_text.setObjectName("Log")
        self.log_text.setReadOnly(True)
        self.log_text.document().setMaximumBlockCount(3000)
        self.log_text.setMinimumHeight(48)
        lay.addWidget(self.log_text)
        return card

    def _toggle_log(self):
        show = not self.log_text.isVisible()
        self.log_text.setVisible(show)
        self.log_toggle.setText("▾ Thu gọn" if show else "▸ Mở rộng")
        if show:
            self.splitter.setSizes([max(self.splitter.height() - 150, 300), 150])
        else:
            self.splitter.setSizes([self.splitter.height() - 50, 50])

    def log(self, message: str):
        now = datetime.now().strftime("%H:%M:%S")
        m = message.lstrip()
        if m.startswith("❌"):
            color = LOG_COLORS["error"]
        elif m.startswith(("✅", "✔", "🎉")):
            color = LOG_COLORS["ok"]
        elif m.startswith(("⚠", "⏳", "⏹")):
            color = LOG_COLORS["warn"]
        elif m.startswith(("💡", "🎙", "🔌", "🧬", "📤", "🔗", "🎬", "🚀", "🔊")):
            color = LOG_COLORS["info"]
        else:
            color = None
        safe = html.escape(message).replace("\n", "<br>")
        body = f"<span style='color:{color}'>{safe}</span>" if color else safe
        self.log_text.append(f"<span style='color:{LOG_COLORS['time']}'>[{now}]</span> {body}")
        sb = self.log_text.verticalScrollBar()
        sb.setValue(sb.maximum())
        self.status_label.setText(message.strip().splitlines()[0][:160] if message.strip() else "")
        try:
            with open(LOG_PATH, "a", encoding="utf-8") as f:
                f.write(f"{datetime.now():%Y-%m-%d %H:%M:%S} {message}\n")
        except Exception:
            pass

    # ================================================================ theme
    def apply_theme(self, mode: str, save: bool = False):
        self.mode = mode
        try:
            arrows = make_arrow_icons(APP_DIR / "ui", mode)
        except Exception:
            arrows = {}
        QApplication.instance().setStyleSheet(build_stylesheet(mode, arrows))
        self.theme_btn.setText("☀  Giao diện sáng" if mode == "dark" else "🌙  Giao diện tối")
        self.theme_btn.setToolTip("Chuyển giao diện Sáng / Tối")
        self.guide.setHtml(self._guide_html())
        self._update_video_color_btn()
        self._update_video_preview()
        self._recolor_tables()
        if save:
            self.config = self.config_store.save({"theme": mode})

    def _recolor_tables(self):
        if hasattr(self, "voice_table"):
            self.refresh_voices()
        for i, st in enumerate(getattr(self, "batch_state", [])):
            self._paint_batch_status(i, st, None)

    # ================================================================ settings
    def _load_settings(self):
        c = self.config
        self.iw_key.setText(self.secret_store.get("inworld_api_key"))
        self.mm_key.setText(self.secret_store.get("minimax_api_key"))
        self.mm_group.setText(c.get("minimax_group_id", ""))
        self.mm_base.setCurrentText(c.get("minimax_base_url", "https://api.minimax.io"))
        self.iw_base.setText(c.get("inworld_base_url", "https://api.inworld.ai"))
        out = c.get("last_output_dir") or str(Path.home() / "TTS_Output")
        self.tts_output_dir.setText(out)
        self.batch_output_dir.setText(c.get("batch_output_dir") or out)
        fmt = c.get("output_format", "mp3")
        self.tts_format.set_value(fmt)
        self.batch_format.set_value(fmt)
        self.batch_threads.setValue(int(c.get("batch_threads", 2) or 2))
        self.batch_skip.setChecked(bool(c.get("batch_skip_existing", True)))
        self.batch_merge.setChecked(bool(c.get("batch_merge_all", False)))
        if c.get("video_size") in VIDEO_SIZES:
            self.video_size.setCurrentText(c["video_size"])
        self.video_color = c.get("video_color", "#101828") or "#101828"
        self._set_video_image(c.get("video_image", ""), save=False)
        self._update_video_color_btn()

    def save_settings(self):
        ok1 = self.secret_store.set("minimax_api_key", self.mm_key.text())
        ok2 = self.secret_store.set("inworld_api_key", self.iw_key.text())
        self.config = self.config_store.save({
            "minimax_group_id": self.mm_group.text().strip(),
            "minimax_base_url": self.mm_base.currentText().strip() or "https://api.minimax.io",
            "inworld_base_url": self.iw_base.text().strip() or "https://api.inworld.ai",
        })
        self._refresh_chips()
        if ok1 and ok2:
            self.log("✅ Đã lưu cài đặt và API key.")
            QMessageBox.information(self, "Đã lưu", "✅ Đã lưu cài đặt.\nAPI key được cất an toàn trong Windows Credential Manager.")
        else:
            self.log("⚠ Đã lưu cài đặt, nhưng không cất được API key vào keyring (chỉ giữ trong phiên này).")
            QMessageBox.warning(self, "Đã lưu một phần",
                                "Cài đặt đã lưu, nhưng keyring không khả dụng trên máy này.\n"
                                "API key chỉ được giữ đến khi đóng app.")

    def _secrets(self):
        return {
            "minimax_api_key": self.mm_key.text().strip() or self.secret_store.get("minimax_api_key"),
            "inworld_api_key": self.iw_key.text().strip() or self.secret_store.get("inworld_api_key"),
        }

    def _live_config(self) -> dict:
        cfg = dict(self.config)
        cfg["minimax_group_id"] = self.mm_group.text().strip()
        cfg["minimax_base_url"] = self.mm_base.currentText().strip() or "https://api.minimax.io"
        cfg["inworld_base_url"] = self.iw_base.text().strip() or "https://api.inworld.ai"
        return cfg

    def _refresh_chips(self):
        s = self._secrets()
        self.chip_inworld.setText(("🟢" if s["inworld_api_key"] else "⚪") + " Inworld")
        self.chip_inworld.setToolTip("Đã có API key Inworld" if s["inworld_api_key"] else "Chưa có API key Inworld")
        self.chip_minimax.setText(("🟢" if s["minimax_api_key"] else "⚪") + " MiniMax")
        self.chip_minimax.setToolTip("Đã có API key MiniMax" if s["minimax_api_key"] else "Chưa có API key MiniMax")
        ff = find_ffmpeg()
        self.chip_ffmpeg.setText(("🟢" if ff else "🔴") + " ffmpeg")
        self.chip_ffmpeg.setToolTip(ff or "Không tìm thấy ffmpeg — không tạo được MP4")
        if hasattr(self, "ffmpeg_label"):
            if ff:
                self.ffmpeg_label.setText(f"✅ Đã sẵn sàng tạo video.\n{ff}")
            else:
                self.ffmpeg_label.setText("❌ Không tìm thấy ffmpeg. Chạy lại install.bat (tự cài imageio-ffmpeg) "
                                          "hoặc chép ffmpeg.exe vào cùng thư mục với ứng dụng.")

    def test_connection(self, provider: str):
        out = self.iw_test_result if provider == "Inworld" else self.mm_test_result
        btn = self.iw_test if provider == "Inworld" else self.mm_test
        out.setText("⏳ Đang kết nối…")
        btn.setEnabled(False)
        secrets, cfg = self._secrets(), self._live_config()

        def job(log, cancelled, progress):
            return build_provider(provider, secrets, cfg).list_voices()

        def ok(voices):
            mine = sum(1 for v in voices if v["kind"] == "Của tôi")
            out.setText(f"✅ Kết nối thành công • {mine} giọng của bạn • {len(voices) - mine} giọng hệ thống.")
            out.setStyleSheet(f"color:{STATUS_COLORS[self.mode]['ok']}")
            self.log(f"✅ {provider}: key hợp lệ ({mine} giọng clone).")

        def bad(err):
            out.setText("❌ " + err)
            out.setStyleSheet(f"color:{STATUS_COLORS[self.mode]['error']}")
            self.log(f"❌ {provider}: kiểm tra kết nối thất bại — {err}")

        self._run(job, ok, bad, finished=lambda: btn.setEnabled(True))

    # ================================================================ helpers
    def _run(self, fn, on_done, on_fail, *, progress=None, finished=None) -> FuncWorker:
        w = FuncWorker(fn, self)
        w.log.connect(self.log)
        w.done.connect(on_done)
        w.failed.connect(on_fail)
        if progress:
            w.progress.connect(progress)

        def _fin():
            self.workers.discard(w)
            if finished:
                finished()

        w.finished.connect(_fin)
        self.workers.add(w)
        w.start()
        return w

    def _pick_folder(self, target: QLineEdit):
        folder = QFileDialog.getExistingDirectory(self, "Chọn thư mục lưu", target.text() or str(Path.home()))
        if folder:
            target.setText(folder)

    def open_folder(self, folder: str):
        if not folder:
            return
        p = Path(folder)
        try:
            p.mkdir(parents=True, exist_ok=True)
        except Exception as exc:
            QMessageBox.warning(self, "Thư mục", f"Không mở được thư mục:\n{exc}")
            return
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(p.resolve())))

    def open_path(self, path: str):
        if path and Path(path).exists():
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(Path(path).resolve())))
        else:
            self.log("⚠ Chưa có file để mở.")

    def _video_settings(self) -> dict:
        return {"image": self.video_image.text().strip(), "size": VIDEO_SIZES[self.video_size.currentText()],
                "color": self.video_color}

    def _format_changed(self, fmt: str):
        for row in (self.tts_format, self.batch_format):
            if row.value() != fmt:
                row.set_value(fmt)
        self.config = self.config_store.save({"output_format": fmt})
        if fmt in ("mp4", "both") and not find_ffmpeg():
            self.log("⚠ Không tìm thấy ffmpeg — sẽ không tạo được MP4. Xem trang 🎬 Video MP4.")

    # ================================================================ voices
    def refresh_voices(self):
        voices = self.voice_store.list()
        sc = STATUS_COLORS[self.mode]
        self.voice_table.setRowCount(len(voices))
        for r, v in enumerate(voices):
            prov = v.get("provider", "")
            vals = [v.get("local_name", ""), prov, v.get("provider_voice_id", ""),
                    language_label(prov, v.get("language")), (v.get("created_at", "") or "").replace("T", " ")[:16]]
            for c, val in enumerate(vals):
                it = QTableWidgetItem(str(val))
                if c == 0:
                    it.setData(Qt.ItemDataRole.UserRole, v.get("uid"))
                    f = it.font()
                    f.setBold(True)
                    it.setFont(f)
                if c == 1:
                    acc = ACCENTS[PROVIDER_ACCENT.get(prov, "blue")]
                    it.setForeground(QColor(acc[1] if self.mode == "dark" else acc[2]))
                    it.setText(("🌊 " if prov == "Inworld" else "🌸 ") + prov)
                    f = it.font()
                    f.setBold(True)
                    it.setFont(f)
                if c == 2:
                    it.setForeground(QColor(sc["pending"]))
                    it.setToolTip(str(val))
                self.voice_table.setItem(r, c, it)
        self.voice_empty.setVisible(not voices)
        self.voice_count_badge.setText(str(len(voices)))
        for combo in (self.tts_voice, self.batch_voice):
            current = combo.currentData()
            combo.blockSignals(True)
            combo.clear()
            if not voices:
                combo.addItem("— Chưa có giọng, hãy Clone hoặc lấy từ tài khoản —", None)
            for v in voices:
                icon = "🌊" if v.get("provider") == "Inworld" else "🌸"
                combo.addItem(f"{icon}  {v.get('local_name')}   ·   {v.get('provider')}", v.get("uid"))
            idx = combo.findData(current) if current else -1
            combo.setCurrentIndex(idx if idx >= 0 else 0)
            combo.blockSignals(False)
        self._voice_changed(self.tts_voice, self.tts_model, self.tts_speed, self.tts_voice_info)
        self._voice_changed(self.batch_voice, self.batch_model, self.batch_speed, self.batch_voice_info)
        if hasattr(self, "chip_voices"):
            self.chip_voices.setText(f"🎤 {len(voices)} giọng")

    def _voice_changed(self, combo, model_combo, speed, info):
        voice = self._voice_from_combo(combo)
        prev_model = model_combo.currentText()
        model_combo.clear()
        if not voice:
            info.setText("")
            speed.setEnabled(False)
            model_combo.setEnabled(False)
            return
        prov = voice.get("provider", "Inworld")
        model_combo.setEnabled(True)
        speed.setEnabled(True)
        model_combo.addItems(MODELS.get(prov, []))
        saved = self.config.get(f"model_{prov}")
        for m in (prev_model, saved):
            if m and m in MODELS.get(prov, []):
                model_combo.setCurrentText(m)
                break
        lo, hi = SPEED_RANGE.get(prov, (0.5, 2.0))
        speed.setRange(lo, hi)
        speed.setValue(float(self.config.get(f"speed_{prov}", 1.0)))
        info.setText(f"🌐 {language_label(prov, voice.get('language'))}   •   ⚡ tốc độ {lo}–{hi}×")
        info.setToolTip("Voice ID: " + voice.get("provider_voice_id", ""))

    def _voice_from_combo(self, combo):
        uid = combo.currentData()
        return self.voice_store.get(uid) if uid else None

    def _selected_voice(self) -> dict | None:
        row = self.voice_table.currentRow()
        if row < 0:
            QMessageBox.information(self, "Chọn giọng", "Hãy bấm chọn một giọng trong bảng trước.")
            return None
        it = self.voice_table.item(row, 0)
        return self.voice_store.get(it.data(Qt.ItemDataRole.UserRole)) if it else None

    def use_selected_voice(self):
        v = self._selected_voice()
        if not v:
            return
        for combo in (self.tts_voice, self.batch_voice):
            idx = combo.findData(v["uid"])
            if idx >= 0:
                combo.setCurrentIndex(idx)
        self.go(PAGE_TTS)
        self.tts_text.setFocus()

    def rename_selected_voice(self):
        v = self._selected_voice()
        if not v:
            return
        name, ok = QInputDialog.getText(self, "Đổi tên giọng", "Tên mới:", text=v.get("local_name", ""))
        if ok and name.strip():
            self.voice_store.update(v["uid"], {"local_name": name.strip()})
            self.refresh_voices()

    def copy_selected_voice_id(self):
        v = self._selected_voice()
        if v:
            QGuiApplication.clipboard().setText(v.get("provider_voice_id", ""))
            self.log(f"📋 Đã sao chép Voice ID: {v.get('provider_voice_id')}")

    def remove_selected_voice(self):
        v = self._selected_voice()
        if not v:
            return
        if QMessageBox.question(
            self, "Xóa giọng",
            f"Xóa “{v.get('local_name')}” khỏi thư viện trong app?\n\n"
            "Giọng trên tài khoản nhà cung cấp KHÔNG bị xóa, có thể lấy lại bằng ☁ Lấy từ tài khoản.",
        ) == QMessageBox.StandardButton.Yes:
            self.voice_store.remove(v["uid"])
            self.refresh_voices()
            self.log(f"🗑 Đã xóa khỏi thư viện: {v.get('local_name')}")

    def _add_voice(self, provider, voice_id, name, language, extra=None) -> bool:
        if self.voice_store.exists(provider, voice_id):
            return False
        now = datetime.now().isoformat(timespec="seconds")
        voice = {"uid": str(uuid.uuid4()), "local_name": name, "provider": provider,
                 "provider_voice_id": voice_id, "language": language, "created_at": now}
        voice.update(extra or {})
        self.voice_store.add(voice)
        return True

    def add_voice_manual(self):
        from dialogs import AddVoiceDialog

        dlg = AddVoiceDialog(self)
        if dlg.exec():
            v = dlg.result_voice()
            if self._add_voice(v["provider"], v["provider_voice_id"], v["local_name"], v["language"]):
                self.refresh_voices()
                self.log(f"✅ Đã thêm giọng: {v['local_name']} ({v['provider']})")
            else:
                QMessageBox.information(self, "Đã có", "Voice ID này đã có trong thư viện.")

    def import_voices(self, provider: str):
        from dialogs import ImportVoicesDialog

        secrets = self._secrets()
        if not secrets["inworld_api_key" if provider == "Inworld" else "minimax_api_key"]:
            self.go(PAGE_SETTINGS)
            QMessageBox.warning(self, "Thiếu API key", f"Hãy nhập API key {provider} ở trang Cài đặt trước.")
            return
        existing = {v.get("provider_voice_id") for v in self.voice_store.list() if v.get("provider") == provider}
        dlg = ImportVoicesDialog(provider, secrets, self._live_config(), existing, self)
        if dlg.exec():
            added = 0
            for v in dlg.selected():
                if self._add_voice(provider, v["voice_id"], v["name"], v.get("language", "")
                                   if provider == "Inworld" else "auto", {"imported": True}):
                    added += 1
            self.refresh_voices()
            self.log(f"✅ Đã thêm {added} giọng từ tài khoản {provider}.")

    def preview_selected_voice(self):
        v = self._selected_voice()
        if not v:
            return
        prov = v["provider"]
        lang = (v.get("language") or "").lower()
        text = PREVIEW_TEXT["en"] if lang.startswith("en") or lang == "english" else PREVIEW_TEXT["vi"]
        out = APP_DIR / "preview" / f"{safe_filename(v.get('local_name', 'voice'))}_{v['uid'][:6]}.mp3"
        out.parent.mkdir(parents=True, exist_ok=True)
        secrets, cfg = self._secrets(), self._live_config()
        model = MODELS[prov][0]
        self.log(f"▶ Đang tạo câu nghe thử cho “{v.get('local_name')}”…")

        def job(log, cancelled, progress):
            build_provider(prov, secrets, cfg, log=log).synthesize(
                text, v["provider_voice_id"], str(out), model=model, language=v.get("language", ""))
            return str(out)

        def ok(path):
            self.log(f"✅ Nghe thử: {path}")
            self.open_path(path)

        self._run(job, ok, lambda e: (self.log("❌ Nghe thử lỗi: " + e),
                                      QMessageBox.critical(self, "Nghe thử thất bại", e)))

    # ================================================================ clone
    def _clone_provider_changed(self, provider):
        self.clone_language.clear()
        for lbl, code in LANGUAGES[provider]:
            self.clone_language.addItem(lbl, code)
        default = "vi-VN" if provider == "Inworld" else "Vietnamese"
        idx = self.clone_language.findData(default)
        if idx >= 0:
            self.clone_language.setCurrentIndex(idx)
        self.clone_rules.setText("ℹ " + sample_hint(provider))
        if self.clone_sample.text().strip():
            self._validate_sample()

    def _validate_sample(self) -> bool:
        path = self.clone_sample.text().strip().strip('"')
        if not path:
            return False
        ok, msg = validate_voice_sample(path, self.clone_provider.value())
        sc = STATUS_COLORS[self.mode]
        color = sc["error"] if not ok else (sc["warn"] if msg.startswith("⚠") else sc["ok"])
        self.clone_sample_info.setText(msg if ok else "❌ " + msg)
        self.clone_sample_info.setStyleSheet(f"color:{color}; font-weight:600;")
        return ok

    def _browse_clone_sample(self):
        path, _ = QFileDialog.getOpenFileName(self, "Chọn file giọng mẫu", "", "Audio (*.wav *.mp3 *.m4a *.webm)")
        if path:
            self.clone_sample.setText(path)
            if not self.clone_name.text().strip():
                self.clone_name.setText(Path(path).stem[:40])
            self._validate_sample()

    def start_clone(self):
        provider = self.clone_provider.value()
        sample = self.clone_sample.text().strip().strip('"')
        name = self.clone_name.text().strip()
        lang = self.clone_language.currentData() or ""
        if not sample:
            QMessageBox.warning(self, "Thiếu file mẫu", "Hãy chọn file giọng mẫu (bước ②).")
            return
        if not self._validate_sample():
            QMessageBox.warning(self, "File mẫu chưa đạt", self.clone_sample_info.text())
            return
        if not name:
            QMessageBox.warning(self, "Thiếu tên", "Hãy đặt tên cho giọng (bước ③).")
            self.clone_name.setFocus()
            return
        if not self.clone_consent.isChecked():
            QMessageBox.warning(self, "Xác nhận quyền sử dụng",
                                "Bạn cần tích ô xác nhận có quyền sử dụng giọng mẫu trước khi clone.")
            return
        secrets = self._secrets()
        if not secrets["inworld_api_key" if provider == "Inworld" else "minimax_api_key"]:
            self.go(PAGE_SETTINGS)
            QMessageBox.warning(self, "Thiếu API key", f"Hãy nhập API key {provider} ở trang Cài đặt.")
            return
        self.config = self.config_store.save({"last_provider": provider})
        cfg = self._live_config()
        denoise = self.clone_denoise.isChecked()
        self.clone_button.setEnabled(False)
        self.clone_button.setText("⏳   Đang clone… (có thể mất 10–60 giây)")
        self.clone_busy.show()
        self.log(f"🧬 Bắt đầu clone “{name}” bằng {provider}…")

        def job(log, cancelled, progress):
            p = build_provider(provider, secrets, cfg, log=log, cancel=cancelled)
            return p.clone_voice(sample, name, lang, denoise=denoise)

        def ok(voice_id):
            self._add_voice(provider, voice_id, name, lang, {
                "sample_file": Path(sample).name,
                "consent_confirmed_at": datetime.now().isoformat(timespec="seconds")})
            self.refresh_voices()
            for combo in (self.tts_voice, self.batch_voice):
                combo.setCurrentIndex(max(0, combo.count() - 1))
            self.voice_table.selectRow(self.voice_table.rowCount() - 1)
            self.log(f"✅ Clone thành công: {name} → {voice_id}")
            extra = ("\n\n⚠ MiniMax sẽ xóa giọng clone nếu 7 ngày không dùng — nên nghe thử ngay."
                     if provider == "MiniMax" else "")
            if QMessageBox.question(self, "🎉 Clone thành công",
                                    f"Đã tạo giọng “{name}”.\nVoice ID: {voice_id}{extra}\n\nNghe thử giọng ngay bây giờ?"
                                    ) == QMessageBox.StandardButton.Yes:
                self.preview_selected_voice()

        def bad(err):
            self.log("❌ Clone thất bại: " + err)
            QMessageBox.critical(self, "Clone thất bại", err)

        def fin():
            self.clone_button.setEnabled(True)
            self.clone_button.setText("🧬   Bắt đầu Clone giọng")
            self.clone_busy.hide()

        self.clone_worker = self._run(job, ok, bad, finished=fin)

    # ================================================================ tts
    def _update_counter(self):
        text = self.tts_text.toPlainText()
        n = len(text.strip())
        voice = self._voice_from_combo(self.tts_voice) if hasattr(self, "tts_voice") else None
        limit = 1800 if not voice or voice.get("provider") == "Inworld" else 5000
        parts = len(split_text(text, limit)) if n else 0
        secs = n / 14.0
        self.tts_counter.setText(f"🔤 {n:,} ký tự   •   🧩 {parts} đoạn gửi API   •   ⏱ ước tính ~{fmt_duration(secs)}"
                                 .replace(",", "."))

    def _paste_text(self):
        t = QGuiApplication.clipboard().text()
        if t:
            self.tts_text.insertPlainText(t)

    def _open_txt(self):
        path, _ = QFileDialog.getOpenFileName(self, "Mở file văn bản", "", "Văn bản (*.txt *.md);;Tất cả (*.*)")
        if not path:
            return
        for enc in ("utf-8-sig", "utf-16", "cp1258", "latin-1"):
            try:
                self.tts_text.setPlainText(Path(path).read_text(encoding=enc))
                if not self.tts_filename.text().strip():
                    self.tts_filename.setText(Path(path).stem)
                return
            except (UnicodeError, ValueError):
                continue

    def start_tts(self):
        voice = self._voice_from_combo(self.tts_voice)
        if not voice:
            QMessageBox.warning(self, "Chưa có giọng", "Hãy clone hoặc thêm một giọng trước (trang 🧬 Clone giọng).")
            return
        text = self.tts_text.toPlainText().strip()
        if not text:
            QMessageBox.warning(self, "Chưa có nội dung", "Hãy nhập nội dung cần đọc.")
            self.tts_text.setFocus()
            return
        outdir = self.tts_output_dir.text().strip()
        if not outdir:
            QMessageBox.warning(self, "Thiếu thư mục", "Hãy chọn thư mục lưu file.")
            return
        try:
            Path(outdir).mkdir(parents=True, exist_ok=True)
        except Exception as exc:
            QMessageBox.critical(self, "Thư mục lỗi", f"Không tạo được thư mục:\n{exc}")
            return
        fmt = self.tts_format.value()
        if fmt in ("mp4", "both") and not find_ffmpeg():
            QMessageBox.critical(self, "Thiếu ffmpeg", "Không tìm thấy ffmpeg nên chưa tạo được MP4. Xem trang 🎬 Video MP4.")
            return
        prov = voice["provider"]
        secrets = self._secrets()
        if not secrets["inworld_api_key" if prov == "Inworld" else "minimax_api_key"]:
            self.go(PAGE_SETTINGS)
            QMessageBox.warning(self, "Thiếu API key", f"Hãy nhập API key {prov} ở trang Cài đặt.")
            return
        stem = self.tts_filename.text().strip() or f"tts_{datetime.now():%Y%m%d_%H%M%S}"
        stem = unique_stem(outdir, stem, exts_for(fmt))
        model = self.tts_model.currentText().strip() or MODELS[prov][0]
        speed = self.tts_speed.value()
        self.config = self.config_store.save({"last_output_dir": outdir, f"model_{prov}": model,
                                              f"speed_{prov}": speed})
        cfg, video = self._live_config(), self._video_settings()

        self.tts_generate.setEnabled(False)
        self.tts_generate.setText("⏳   Đang tạo…")
        self.tts_stop.setEnabled(True)
        self.tts_progress.setValue(0)
        self.tts_result.setText("")
        self.log(f"🔊 Tạo giọng đọc “{stem}” bằng {voice.get('local_name')} ({prov}, {model}, {speed:.2f}×) — {len(text)} ký tự")

        def job(log, cancelled, progress):
            p = build_provider(prov, secrets, cfg, log=log, cancel=cancelled)
            return produce_outputs(provider=p, text=text, voice=voice, model=model, speed=speed, out_dir=outdir,
                                   name=stem, fmt=fmt, video=video, log=log, cancelled=cancelled, progress=progress)

        def ok(res):
            files = [p for p in (res["mp3"], res["mp4"]) if p]
            self.last_outputs = files
            dur = probe_duration(res["audio"])
            names = "  +  ".join(Path(f).name for f in files)
            self.tts_result.setText(f"✅ Đã tạo: {names}   •   ⏱ {fmt_duration(dur)}\n📁 {outdir}")
            self.tts_result.setStyleSheet(f"color:{STATUS_COLORS[self.mode]['ok']}; font-weight:600;")
            self.tts_open_file.setEnabled(True)
            self.tts_progress.setValue(100)
            self.log(f"✅ Đã tạo: {names} ({fmt_duration(dur)})")

        def bad(err):
            self.tts_progress.setValue(0)
            self.tts_result.setText("❌ " + err)
            self.tts_result.setStyleSheet(f"color:{STATUS_COLORS[self.mode]['error']}; font-weight:600;")
            self.log("❌ Tạo giọng đọc lỗi: " + err)
            if "Đã dừng" not in err:
                QMessageBox.critical(self, "Tạo giọng đọc thất bại", err)

        def fin():
            self.tts_generate.setEnabled(True)
            self.tts_generate.setText("🔊   Tạo giọng đọc")
            self.tts_stop.setEnabled(False)

        self.tts_worker = self._run(job, ok, bad, progress=self.tts_progress.setValue, finished=fin)

    def stop_tts(self):
        if self.tts_worker and self.tts_worker.isRunning():
            self.tts_worker.cancel()
            self.tts_stop.setEnabled(False)
            self.log("⏹ Đang dừng…")

    def _open_last_output(self):
        if self.last_outputs:
            self.open_path(self.last_outputs[-1])

    # ================================================================ batch
    def select_batch_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Chọn file Excel", "", "Excel (*.xlsx *.xlsm)")
        if path:
            self._open_batch_file(path)

    def _reload_batch_file(self):
        if self.batch_path:
            self._open_batch_file(self.batch_path, keep_sheet=True)

    def _open_batch_file(self, path: str, keep_sheet: bool = False):
        from openpyxl import load_workbook

        prev_sheet = self.batch_sheet.currentText()
        self.batch_path = path
        self.batch_file.setText(path)
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            names = list(wb.sheetnames)
            wb.close()
        except PermissionError:
            QMessageBox.critical(self, "Không đọc được Excel", "File đang được mở/khóa bởi chương trình khác. Hãy đóng Excel rồi thử lại.")
            return
        except Exception as exc:
            QMessageBox.critical(self, "Không đọc được Excel", str(exc))
            return
        self.batch_sheet.blockSignals(True)
        self.batch_sheet.clear()
        self.batch_sheet.addItems(names)
        if keep_sheet and prev_sheet in names:
            self.batch_sheet.setCurrentText(prev_sheet)
        self.batch_sheet.blockSignals(False)
        self._batch_sheet_changed(self.batch_sheet.currentText())
        self.log(f"📗 Đã mở Excel: {Path(path).name} ({len(names)} sheet)")

    def _batch_sheet_changed(self, sheet):
        from openpyxl import load_workbook

        for c in (self.batch_text_col, self.batch_name_col):
            c.blockSignals(True)
            c.clear()
        try:
            if not self.batch_path or not sheet:
                return
            wb = load_workbook(self.batch_path, read_only=True, data_only=True)
            ws = wb[sheet]
            first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            wb.close()
            headers = [str(v).strip() if v is not None else "" for v in first]
            self.batch_name_col.addItem("🔢 Tự đánh số (0001, 0002, …)", 0)
            for idx, h in enumerate(headers, 1):
                col = chr(64 + idx) if idx <= 26 else f"#{idx}"
                lbl = f"{col}: {h}" if h else f"Cột {col}"
                self.batch_text_col.addItem(lbl, idx)
                self.batch_name_col.addItem(lbl, idx)
            lower = [h.lower() for h in headers]

            def pick(combo, cands, offset):
                for cand in cands:
                    for i, h in enumerate(lower):
                        if h == cand:
                            combo.setCurrentIndex(i + offset)
                            return True
                for cand in cands:
                    for i, h in enumerate(lower):
                        if cand in h:
                            combo.setCurrentIndex(i + offset)
                            return True
                return False

            if not pick(self.batch_text_col, ("text", "nội dung", "noi dung", "content", "văn bản", "van ban",
                                              "script", "kịch bản", "tts", "lời thoại"), 0):
                # choose the column with the longest header-free guess: last column
                self.batch_text_col.setCurrentIndex(max(0, len(headers) - 1))
            if not pick(self.batch_name_col, ("filename", "file name", "tên file", "ten file", "file", "stt",
                                              "id", "name", "tên", "title"), 1):
                self.batch_name_col.setCurrentIndex(0)
            if self.batch_name_col.currentData() == self.batch_text_col.currentData():
                self.batch_name_col.setCurrentIndex(0)
        except Exception as exc:
            self.log(f"❌ Lỗi đọc tiêu đề Excel: {exc}")
        finally:
            for c in (self.batch_text_col, self.batch_name_col):
                c.blockSignals(False)
        self._batch_load_tasks()

    def _read_batch_tasks(self) -> list[dict]:
        from openpyxl import load_workbook

        if not self.batch_path:
            raise RuntimeError("Chưa chọn file Excel.")
        text_col = self.batch_text_col.currentData()
        name_col = self.batch_name_col.currentData()
        if not text_col:
            raise RuntimeError("Chưa chọn cột nội dung.")
        wb = load_workbook(self.batch_path, read_only=True, data_only=True)
        try:
            ws = wb[self.batch_sheet.currentText()]
            tasks, used = [], set()
            for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                tv = row[text_col - 1] if text_col - 1 < len(row) else None
                text = str(tv).strip() if tv is not None else ""
                if not text:
                    continue
                if name_col:
                    nv = row[name_col - 1] if name_col - 1 < len(row) else None
                    base = safe_filename(nv, f"{len(tasks) + 1:04d}") if nv not in (None, "") else f"{len(tasks) + 1:04d}"
                else:
                    base = f"{len(tasks) + 1:04d}"
                name, i = base, 2
                while name.lower() in used:
                    name = f"{base}_{i}"
                    i += 1
                used.add(name.lower())
                tasks.append({"row_no": row_no, "text": text, "filename": name})
            return tasks
        finally:
            wb.close()

    def _batch_load_tasks(self):
        if self.batch_worker and self.batch_worker.isRunning():
            return
        try:
            self.batch_tasks = self._read_batch_tasks() if self.batch_path else []
        except Exception as exc:
            self.batch_tasks = []
            self.log(f"❌ Lỗi đọc Excel: {exc}")
        self.batch_state = ["pending"] * len(self.batch_tasks)
        t = self.batch_table
        t.setRowCount(len(self.batch_tasks))
        for i, task in enumerate(self.batch_tasks):
            preview = task["text"].replace("\n", " ")
            vals = [str(task["row_no"]), task["filename"], preview[:160] + ("…" if len(preview) > 160 else ""),
                    f"{len(task['text']):,}".replace(",", "."), ""]
            for c, val in enumerate(vals):
                it = QTableWidgetItem(val)
                if c == 2:
                    it.setToolTip(task["text"][:1500])
                if c in (0, 3):
                    it.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                t.setItem(i, c, it)
            self._paint_batch_status(i, "pending", "")
        self._update_tiles()
        self.batch_progress.setValue(0)
        self.batch_retry.setEnabled(False)
        if self.batch_path:
            chars = sum(len(x["text"]) for x in self.batch_tasks)
            self.batch_status.setText(f"Sẵn sàng: {len(self.batch_tasks)} dòng có nội dung • ~{fmt_duration(chars / 14)} audio.")

    def _paint_batch_status(self, i: int, state: str, detail: str | None):
        if i >= self.batch_table.rowCount():
            return
        it = self.batch_table.item(i, 4)
        if it is None:
            it = QTableWidgetItem()
            self.batch_table.setItem(i, 4, it)
        if detail is not None:
            txt = STATUS_TEXT.get(state, state) + (f"  {detail}" if detail and state != "pending" else "")
            it.setText(txt)
            it.setToolTip(detail or "")
        it.setForeground(QColor(STATUS_COLORS[self.mode].get(state, STATUS_COLORS[self.mode]["pending"])))
        f = it.font()
        f.setBold(state in ("ok", "error", "running"))
        it.setFont(f)

    def _update_tiles(self):
        st = self.batch_state
        self.tile_total.set(len(st))
        self.tile_ok.set(st.count("ok"))
        self.tile_err.set(st.count("error"))
        self.tile_skip.set(st.count("skipped"))
        chars = sum(len(t["text"]) for t in self.batch_tasks)
        self.tile_chars.set(f"{chars:,}".replace(",", "."))

    def create_sample_excel(self):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        path, _ = QFileDialog.getSaveFileName(self, "Lưu file Excel mẫu",
                                              str(Path(self.batch_output_dir.text() or Path.home()) / "mau_tts.xlsx"),
                                              "Excel (*.xlsx)")
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "TTS"
        ws.append(["filename", "text"])
        rows = [
            ("0001", "Xin chào các bạn, chào mừng đến với kênh của mình."),
            ("0002", "Hôm nay chúng ta sẽ cùng tìm hiểu cách tạo giọng đọc tự động từ file Excel."),
            ("0003", "Mỗi dòng trong file này sẽ trở thành một file âm thanh hoặc video riêng."),
        ]
        for r in rows:
            ws.append(list(r))
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="10B981")
            c.alignment = Alignment(horizontal="center")
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 90
        for row in ws.iter_rows(min_row=2):
            row[0].number_format = "@"
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
        try:
            wb.save(path)
        except PermissionError:
            QMessageBox.critical(self, "Không lưu được", "File đang mở trong Excel. Hãy đóng lại rồi thử lại.")
            return
        self.log(f"📥 Đã tạo file mẫu: {path}")
        self._open_batch_file(path)

    def start_batch(self, only_failed: bool = False):
        voice = self._voice_from_combo(self.batch_voice)
        if not voice:
            QMessageBox.warning(self, "Chưa có giọng", "Hãy chọn giọng cho chạy hàng loạt.")
            return
        if not only_failed:
            self._batch_load_tasks()
        if not self.batch_tasks:
            QMessageBox.warning(self, "Chưa có dữ liệu", "Hãy chọn file Excel có ít nhất 1 dòng nội dung.")
            return
        indices = [i for i, s in enumerate(self.batch_state) if s in ("error", "stopped")] if only_failed \
            else list(range(len(self.batch_tasks)))
        if not indices:
            QMessageBox.information(self, "Không có dòng lỗi", "Không có dòng nào cần chạy lại.")
            return
        out_dir = self.batch_output_dir.text().strip()
        if not out_dir:
            QMessageBox.warning(self, "Thiếu thư mục", "Hãy chọn thư mục lưu file.")
            return
        try:
            Path(out_dir).mkdir(parents=True, exist_ok=True)
        except Exception as exc:
            QMessageBox.critical(self, "Thư mục lỗi", f"Không tạo được thư mục:\n{exc}")
            return
        fmt = self.batch_format.value()
        if fmt in ("mp4", "both") and not find_ffmpeg():
            QMessageBox.critical(self, "Thiếu ffmpeg", "Không tìm thấy ffmpeg nên chưa tạo được MP4. Xem trang 🎬 Video MP4.")
            return
        prov = voice["provider"]
        secrets = self._secrets()
        if not secrets["inworld_api_key" if prov == "Inworld" else "minimax_api_key"]:
            self.go(PAGE_SETTINGS)
            QMessageBox.warning(self, "Thiếu API key", f"Hãy nhập API key {prov} ở trang Cài đặt.")
            return
        model = self.batch_model.currentText().strip() or MODELS[prov][0]
        speed = self.batch_speed.value()
        self.config = self.config_store.save({
            "batch_output_dir": out_dir, "batch_threads": self.batch_threads.value(),
            "batch_skip_existing": self.batch_skip.isChecked(), "batch_merge_all": self.batch_merge.isChecked(),
            f"model_{prov}": model, f"speed_{prov}": speed,
        })
        self.batch_map = indices
        for i in indices:
            self.batch_state[i] = "pending"
            self._paint_batch_status(i, "pending", "")
        self._update_tiles()
        merge = self.batch_merge.isChecked() and not only_failed
        self.batch_progress.setRange(0, len(indices))
        self.batch_progress.setValue(0)
        self.batch_status.setText(f"Đang chạy 0/{len(indices)}…")
        self.batch_run.setEnabled(False)
        self.batch_retry.setEnabled(False)
        self.batch_stop.setEnabled(True)
        self.batch_run.setText("⏳   Đang chạy…")
        self.log(f"🚀 Chạy {'lại ' if only_failed else ''}{len(indices)} dòng bằng {voice.get('local_name')} "
                 f"({prov}, {model}, {speed:.2f}×, {self.batch_threads.value()} luồng, định dạng {fmt.upper()})")

        self.batch_worker = BatchWorker(
            tasks=[self.batch_tasks[i] for i in indices], voice=voice, model=model, speed=speed, out_dir=out_dir,
            fmt=fmt, video=self._video_settings(), threads=self.batch_threads.value(),
            skip_existing=self.batch_skip.isChecked(), merge_all=merge,
            merge_name=safe_filename(Path(self.batch_path).stem + "_GOP"),
            secrets=secrets, config=self._live_config(), parent=self,
        )
        self.batch_worker.log.connect(self.log)
        self.batch_worker.row_status.connect(self._batch_row_status)
        self.batch_worker.progress.connect(self._batch_progress)
        self.batch_worker.finished_summary.connect(self._batch_done)
        self.batch_worker.finished.connect(self._batch_finished)
        self.workers.add(self.batch_worker)
        self.batch_worker.start()

    def _batch_row_status(self, sub_idx: int, state: str, detail: str):
        if sub_idx >= len(self.batch_map):
            return
        i = self.batch_map[sub_idx]
        self.batch_state[i] = state
        self._paint_batch_status(i, state, detail)
        if state == "running":
            self.batch_table.scrollToItem(self.batch_table.item(i, 0))
        self._update_tiles()

    def _batch_progress(self, done: int, total: int):
        self.batch_progress.setRange(0, max(total, 1))
        self.batch_progress.setValue(done)
        self.batch_status.setText(f"Đang chạy {done}/{total}…")

    def _batch_done(self, s: dict):
        msg = f"Thành công {s.get('ok', 0)} • Lỗi {s.get('error', 0)} • Bỏ qua {s.get('skipped', 0)}"
        if s.get("stopped"):
            msg += f" • Dừng {s['stopped']}"
        head = "⏹ Đã dừng" if s.get("cancelled") else "🎉 Hoàn tất"
        self.batch_status.setText(f"{head}: {msg}")
        self.log(("⏹ " if s.get("cancelled") else "✅ ") + f"Batch xong — {msg}")
        if s.get("merged"):
            self.last_outputs = s["merged"]
        if s.get("fatal"):
            QMessageBox.critical(self, "Batch lỗi", s["fatal"])
        elif not s.get("cancelled"):
            extra = ("\n\n🔗 File gộp:\n" + "\n".join(Path(m).name for m in s["merged"])) if s.get("merged") else ""
            box = QMessageBox(QMessageBox.Icon.Information, "Hoàn tất hàng loạt",
                              f"✅ Thành công: {s.get('ok', 0)}\n❌ Lỗi: {s.get('error', 0)}\n"
                              f"⏭ Bỏ qua: {s.get('skipped', 0)}{extra}", parent=self)
            open_btn = box.addButton("📂 Mở thư mục", QMessageBox.ButtonRole.AcceptRole)
            box.addButton("Đóng", QMessageBox.ButtonRole.RejectRole)
            box.exec()
            if box.clickedButton() is open_btn:
                self.open_folder(self.batch_output_dir.text())

    def _batch_finished(self):
        self.workers.discard(self.batch_worker)
        self.batch_run.setEnabled(True)
        self.batch_run.setText("🚀   Chạy hàng loạt")
        self.batch_stop.setEnabled(False)
        self.batch_retry.setEnabled(any(s in ("error", "stopped") for s in self.batch_state))

    def stop_batch(self):
        if self.batch_worker and self.batch_worker.isRunning():
            self.batch_worker.cancel()
            self.batch_stop.setEnabled(False)
            self.batch_status.setText("⏹ Đang dừng — chờ các dòng đang chạy kết thúc…")
            self.log("⏹ Đã yêu cầu dừng batch.")

    # ================================================================ video
    def _pick_video_image(self):
        path, _ = QFileDialog.getOpenFileName(self, "Chọn ảnh nền", "", "Ảnh (*.jpg *.jpeg *.png *.webp *.bmp)")
        if path:
            self._set_video_image(path)

    def _set_video_image(self, path: str, save: bool = True):
        if path and not Path(path).is_file():
            path = ""
        self.video_image.setText(path)
        self._update_video_preview()
        if save:
            self.config = self.config_store.save({"video_image": path})

    def _video_size_changed(self, key: str):
        if key in VIDEO_SIZES and hasattr(self, "config_store"):
            self.config = self.config_store.save({"video_size": key})
        self._update_video_preview()

    def _pick_video_color(self):
        c = QColorDialog.getColor(QColor(self.video_color), self, "Chọn màu nền")
        if c.isValid():
            self.video_color = c.name()
            self._update_video_color_btn()
            self._update_video_preview()
            self.config = self.config_store.save({"video_color": self.video_color})

    def _update_video_color_btn(self):
        if not hasattr(self, "video_color_btn"):
            return
        pm = QPixmap(18, 18)
        pm.fill(QColor(self.video_color))
        from PyQt6.QtGui import QIcon

        self.video_color_btn.setIcon(QIcon(pm))
        self.video_color_btn.setText(f"  {self.video_color.upper()}   •   Đổi màu")

    def _update_video_preview(self):
        if not hasattr(self, "video_preview"):
            return
        size_key = self.video_size.currentText()
        w, h = VIDEO_SIZES.get(size_key, (1280, 720))
        box_w = max(self.video_preview.width() - 20, 400)
        box_h = max(self.video_preview.height() - 20, 280)
        scale = min(box_w / w, box_h / h)
        pw, ph = int(w * scale), int(h * scale)
        pm = QPixmap(pw, ph)
        pm.fill(QColor(self.video_color))
        painter = QPainter(pm)
        painter.setRenderHint(QPainter.RenderHint.SmoothPixmapTransform)
        img_path = self.video_image.text().strip()
        if img_path:
            src = QPixmap(img_path)
            if not src.isNull():
                s = src.scaled(pw, ph, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation)
                painter.drawPixmap((pw - s.width()) // 2, (ph - s.height()) // 2, s)
        else:
            painter.setPen(QColor(255, 255, 255, 170))
            f = QFont(self.font())
            f.setPointSize(max(10, int(ph / 14)))
            f.setBold(True)
            painter.setFont(f)
            painter.drawText(QRectF(0, 0, pw, ph), Qt.AlignmentFlag.AlignCenter, "🎙  Màu nền")
        painter.end()
        self.video_preview.setPixmap(pm)
        self.video_preview_info.setText(f"{w}×{h} px  •  H.264 + AAC  •  "
                                        + ("ảnh: " + Path(img_path).name if img_path else "không dùng ảnh"))

    def make_test_video(self):
        ff = find_ffmpeg()
        if not ff:
            QMessageBox.critical(self, "Thiếu ffmpeg", "Không tìm thấy ffmpeg.")
            return
        outdir = self.tts_output_dir.text().strip() or str(Path.home())
        video = self._video_settings()
        self.video_test_btn.setEnabled(False)
        self.log("🎬 Đang tạo video thử 3 giây…")

        def job(log, cancelled, progress):
            import subprocess
            import tempfile

            with tempfile.TemporaryDirectory() as td:
                a = str(Path(td) / "tone.mp3")
                subprocess.run([ff, "-y", "-hide_banner", "-loglevel", "error", "-f", "lavfi", "-i",
                                "sine=frequency=440:duration=3", "-c:a", "libmp3lame", a],
                               capture_output=True, check=True,
                               creationflags=0x08000000 if os.name == "nt" else 0)
                Path(outdir).mkdir(parents=True, exist_ok=True)
                out = str(Path(outdir) / "video_thu.mp4")
                make_video(a, out, image_path=video["image"], size=video["size"], bg_color=video["color"])
                return out

        def ok(path):
            self.log(f"✅ Video thử: {path}")
            self.open_path(path)

        self._run(job, ok, lambda e: (self.log("❌ Video thử lỗi: " + e), QMessageBox.critical(self, "Lỗi", e)),
                  finished=lambda: self.video_test_btn.setEnabled(True))

    # ================================================================ window
    def resizeEvent(self, e):
        super().resizeEvent(e)
        QTimer.singleShot(0, self._apply_responsive)

    def showEvent(self, e):
        super().showEvent(e)
        QTimer.singleShot(0, self._apply_responsive)
        if not getattr(self, "_shown_once", False):
            self._shown_once = True
            if self.height() < 700:  # small / high-DPI laptop screens: give pages the room
                QTimer.singleShot(0, self._toggle_log)

    def _apply_responsive(self):
        if not hasattr(self, "stack"):
            return
        avail = self.stack.width() - 24
        for r in self.responsive:
            r.update_for(avail)
        room = self.chip_box.width() - 10
        for c in (self.chip_ffmpeg, self.chip_minimax, self.chip_inworld, self.chip_voices):
            need = c.sizeHint().width() + 8
            c.setVisible(room >= need)
            if room >= need:
                room -= need
        if self.stack.currentIndex() == PAGE_VIDEO:
            self._update_video_preview()

    def _restore_geometry(self):
        geo = self.config.get("window_geometry")
        if geo:
            try:
                if self.restoreGeometry(QByteArray.fromBase64(geo.encode("ascii"))):
                    return
            except Exception:
                pass
        scr = QGuiApplication.primaryScreen().availableGeometry()
        w, h = min(1400, scr.width() - 40), min(900, scr.height() - 60)
        self.resize(w, h)
        self.move(scr.x() + (scr.width() - w) // 2, scr.y() + (scr.height() - h) // 2)

    def closeEvent(self, e):
        running = [w for w in list(self.workers) if w.isRunning()]
        if running and not self._quitting_for_update:
            if QMessageBox.question(self, "Đang xử lý",
                                    "Vẫn còn tác vụ đang chạy. Dừng lại và thoát?") != QMessageBox.StandardButton.Yes:
                e.ignore()
                return
            for w in running:
                w.cancel()
            for w in running:
                w.wait(8000)
        try:
            self.config_store.save({"window_geometry": bytes(self.saveGeometry().toBase64()).decode("ascii")})
        except Exception:
            pass
        e.accept()


def _install_excepthook():
    def hook(exc_type, exc, tb):
        text = "".join(traceback.format_exception(exc_type, exc, tb))
        try:
            APP_DIR.mkdir(parents=True, exist_ok=True)
            with open(LOG_PATH, "a", encoding="utf-8") as f:
                f.write(f"\n==== {datetime.now():%Y-%m-%d %H:%M:%S} UNCAUGHT ====\n{text}\n")
        except Exception:
            pass
        if os.getenv("TTS_SELFCHECK"):
            sys.__excepthook__(exc_type, exc, tb)
            return
        try:
            QMessageBox.critical(None, "Lỗi không mong muốn",
                                 f"{exc}\n\nChi tiết đã ghi vào:\n{LOG_PATH}")
        except Exception:
            sys.__excepthook__(exc_type, exc, tb)

    sys.excepthook = hook


def main():
    if os.name == "nt":
        try:
            import ctypes

            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("TTSCloneStudio.App")
        except Exception:
            pass
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    app.setOrganizationName("TTSCloneStudio")
    app.setStyle("Fusion")
    from PyQt6.QtGui import QIcon

    app.setWindowIcon(QIcon(resource_path("assets/app.ico" if os.name == "nt" else "assets/app.png")))
    _install_excepthook()
    w = MainWindow()
    w.show()
    if os.getenv("TTS_SELFCHECK"):  # automated check of a built app (CI)
        import selfcheck

        QTimer.singleShot(1500, lambda: selfcheck.run(w, app))
        QTimer.singleShot(240000, app.quit)  # hard stop
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
